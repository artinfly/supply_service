"""Сводки, выгрузки в Excel, краткая справка и анализ ГОЗ."""

import io
import zipfile
from datetime import datetime
from decimal import Decimal

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse

from core.models import GozContractVat
from core.services.goz import analysis
from core.services.goz.gk_directory import (
    directory_index,
    find_gk,
    import_extended_rows,
    import_simple_values,
    simple_values_from_rows,
)

from .helpers import (
    contract_row,
    contracts_file,
    make_user,
    messages_text,
    upload,
    znp_file,
    znp_row,
)


def workbook(response):
    """Первый лист выгрузки как список строк."""
    book = openpyxl.load_workbook(io.BytesIO(response.content))
    return [list(r) for r in book.active.iter_rows(values_only=True)]


class ContractReportsTests(TestCase):
    """Две позиции заключённого договора, черновик и расторгнутый договор (2026 год)."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = make_user("admin", superuser=True)

    def setUp(self):
        self.client.force_login(self.admin)
        upload(
            self.client,
            "contracts",
            contracts_file(
                contract_row(plan=1_000_000, fact=400_000),
                contract_row(order="1002", stage="Этап 2", plan=500_000, fact=0),
                contract_row(
                    contract="Д-2/25", order="1003", status="Черновик", plan=2_000_000
                ),
                contract_row(contract="Д-3/25", order="1004", status="Расторгнут"),
            ),
        )
        upload(
            self.client, "znp", znp_file(znp_row(plan_sum=400_000, fact_sum=400_000))
        )

    def test_dashboard_cards(self):
        ctx = self.client.get(reverse("dashboard"), {"year": 2026}).context
        self.assertEqual(ctx["all_contracts_count"], 3)  # «всего» включает расторгнутые
        self.assertEqual(ctx["curr_year_contracts_count"], 2)
        self.assertEqual(ctx["curr_year_concluded_count"], 1)
        self.assertEqual(ctx["curr_year_not_concluded_count"], 1)
        # Аванс по плану за год — все позиции «Аванс» без расторгнутых, в млн ₽
        self.assertEqual(ctx["curr_year_plan"], Decimal("3.5"))
        self.assertEqual(ctx["igk_table"][-1]["cfo"], "ИТОГО")

    def test_znp_and_sap_pages(self):
        ctx = self.client.get(reverse("znp_table"), {"year": 2026}).context
        self.assertEqual(ctx["all"]["total_count"], ctx["year"]["total_count"])
        self.assertTrue(ctx["has_data"])
        sap = self.client.get(reverse("znp_sap_table"), {"date": "2026-08-21"})
        self.assertEqual(sap.status_code, 200)
        self.assertFalse(sap.context["has_data"])

    def test_registry_api(self):
        rows = self.client.get(
            reverse("api_all_contracts"), {"status": ["Подписан"], "year": 2026}
        ).json()
        positions = [r for r in rows if r.get("is_subtotal") != 1]
        self.assertEqual(sorted(r["order"] for r in positions), ["1001", "1002"])
        self.assertEqual(
            self.client.get(reverse("api_all_contracts"), {"status": ""}).json(), []
        )

    def test_kdr_export(self):
        rows = workbook(self.client.get(reverse("export_kdr", args=[2026])))
        self.assertIn("Всего договоров на 2026г., шт.", rows[0])
        self.assertEqual(rows[0][12], "Заключено, шт.")
        self.assertEqual([r[0] for r in rows[1:]], ["000000001111", None, "ИТОГО"])
        total = rows[-1]
        self.assertEqual(total[6], 2)  # договоров за год без расторгнутых
        self.assertEqual(total[8], 1)  # из них заключено
        self.assertEqual(total[9], 50.0)  # % по количеству

    def test_kdr_export_with_period(self):
        response = self.client.get(
            reverse("export_kdr", args=[2026]),
            {"start": "2026-01-01", "end": "2026-12-31"},
        )
        self.assertEqual(
            workbook(response)[0][12], "Заключено с 01.01.2026 по 31.12.2026, шт."
        )

    def test_export_rejects_bad_params(self):
        self.assertEqual(
            self.client.get(reverse("export_kdr", args=[1999])).status_code, 400
        )
        response = self.client.get(
            reverse("export_kdr", args=[2026]),
            {"start": "2026-13-01", "end": "2026-12-31"},
        )
        self.assertEqual(response.json(), {"error": "недопустимая дата периода"})

    def test_other_exports_open(self):
        urls = [
            reverse("export_advances", args=[2026]),
            reverse("export_contracts_by_agent", args=[2026]),
            reverse("export_history_status"),
            reverse("export_contract_dupes"),
            reverse("export_appeared_concluded"),
        ]
        for url in urls:
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 200)
                self.assertTrue(workbook(response))


def sum_report_file():
    """«Свод отчёт по ИГК»: лист на каждый ИГК, данные с третьей строки."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, column_g in (("1111", "Длинноцикловые"), ("2222", "Этап")):
        ws = wb.create_sheet(title)
        ws["B1"] = f"Свод отчёт по ИГК {title} на 2026 год"
        ws.cell(2, 7, column_g)
        for i, status in enumerate(
            ["Заключен", "Не заключен", "Заключен", "Расторгнут"]
        ):
            values = {1: title, 2: f"ООО Контрагент {i}", 4: "421", 5: f"Д-{i}", 6: status,
                      7: "Да" if i == 0 else "Нет", 8: f"Предмет {i}", 10: 2_000_000.0 * (i + 1),
                      11: 100_000.0, 14: 50_000.0, 18: datetime(2026, 5, 1), 21: "15.06.2026",
                      24: 1000.0, 26: 2000.0, 27: 1, 28: 800.0, 29: 300.0, 32: 700.0, 33: 400.0}  # fmt: skip
            for col, value in values.items():
                ws.cell(3 + i, col, value)
        ws.cell(7, 1, "Итого по ГК")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class SumReportTests(TestCase):
    def setUp(self):
        self.client.force_login(make_user("viewer", "access_export"))

    def post(self, content, kind="normal"):
        return self.client.post(
            reverse("sum_report_upload"),
            {
                "report_type": kind,
                "excel_file": SimpleUploadedFile("svod.xlsx", content),
            },
        )

    def test_one_report_per_igk_sheet(self):
        for kind in ("normal", "cycle"):
            with self.subTest(kind=kind):
                response = self.post(sum_report_file(), kind)
                self.assertEqual(response["Content-Type"], "application/zip")
                names = zipfile.ZipFile(io.BytesIO(response.content)).namelist()
                self.assertEqual(
                    sorted(n.split()[0] for n in names),
                    ["1111", "2222"] if kind == "normal" else ["1111"],
                )
                self.assertTrue(all(n.endswith("_viewer.xlsx") for n in names))

    def test_cycle_report_needs_cycle_column(self):
        wb = openpyxl.load_workbook(io.BytesIO(sum_report_file()))
        wb["1111"].cell(2, 7, "Этап")
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertLogs("core.views.sum_report", "ERROR"):
            response = self.post(buf.getvalue(), "cycle")
        self.assertIn("нет колонки «Длинноцикловые»", response.json()["error"])

    def test_errors_come_back_as_json(self):
        self.assertEqual(
            self.client.post(reverse("sum_report_upload")).json(),
            {"error": "Файл не прикреплён"},
        )
        with self.assertLogs("core.views.sum_report", "ERROR"):
            response = self.post(b"not an excel file")
        self.assertEqual(response.status_code, 400)
        self.assertIn("error", response.json())


class GozTests(TestCase):
    PLAN = {
        "shipment": 1200,
        "cost": 800,
        "amr": 50,
        "commercial": 20,
        "credit_pct": 10,
        "profit": 120,
    }

    def test_build_report(self):
        fact = {**self.PLAN, "profit": 150, "financing": 1000, "distribution": 0, "materials": 0,
                "vat_in": 0, "wip": 0, "resource_delta": 0}  # fmt: skip
        data = analysis.build_report(
            [("1493", self.PLAN, fact)], {"1493": "20"}, {"1493": "Изделие"}
        )
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertEqual(ws["A1"].value, "Анализ отчётов ЕИС ГОЗ")
        self.assertEqual(
            [ws.cell(r, 2).value for r in (3, 4)], ["1493", None]
        )  # ГК объединена на блок
        self.assertEqual(
            [ws.cell(r, 6).value for r in range(3, 7)][:2], ["Целевые", "Факт"]
        )

    def test_directory_lookup(self):
        # Номер ГК, записанный в Excel числом, теряет ведущие нули — они дописываются
        values = simple_values_from_rows([["№", "ГК"], [1, 509], [2, "1493"]])
        self.assertEqual(values, ["0509", "1493"])
        import_simple_values([*values, "0123450262/255сб/137 от 01.02.2024, спец/136"])
        self.assertEqual(sorted(directory_index()), ["0262/255сб 136", "0509", "1493"])
        GozContractVat.objects.filter(igk="0262/255сб 136").update(
            contract_date="Д-1 от 01.02.2024"
        )
        # В имени файла спецификации номер сб не пишется — запись находится по ИГК и спец
        self.assertEqual(find_gk("0262_136", directory_index()).igk, "0262/255сб 136")
        self.assertIsNone(find_gk("9999", directory_index()))
        self.assertEqual(import_simple_values(["1493"]), (0, 1, []))

    def test_profitability_uses_full_cost(self):
        # Рентабельность = прибыль / (себестоимость + АУР + коммерческие)
        metrics = analysis._line_metrics(self.PLAN, 20.0)
        self.assertAlmostEqual(metrics[9], 120 / (800 + 50 + 20))

    def test_same_ending_in_directory(self):
        added, _updated, errors = import_extended_rows(
            [
                ["№", "Предприятие", "Заказ", "Договор/Дата"],
                [1, "АО А", "З-1", "2123000195369/137 от 01.02.2024"],
                [2, "АО Б", "З-2", "2527000195369/137 от 01.02.2024"],
            ]
        )
        self.assertEqual((added, errors), (2, []))
        self.assertEqual(sorted(directory_index()), ["2123_5369", "2527_5369"])

    def test_prefixed_archive_key_finds_plain_record(self):
        # Старая запись без префикса подходит, только если её номер начинается с префикса
        GozContractVat.objects.create(
            igk="5369", contract_date="2123000195369/137 от 01.02.2024"
        )
        index = directory_index()
        self.assertEqual(find_gk("2123_5369", index).igk, "5369")
        self.assertIsNone(find_gk("2527_5369", index))

    def test_archive_errors(self):
        self.client.force_login(make_user("goz", "access_goz_report"))
        not_zip = SimpleUploadedFile("a.zip", b"plain text")
        response = self.client.post(
            reverse("goz_report"), {"archive": not_zip}, follow=True
        )
        self.assertEqual(messages_text(response), ["Файл не является ZIP-архивом"])

        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as zf:
            zf.writestr("readme.txt", "нет отчётов")
        empty = SimpleUploadedFile("a.zip", buf.getvalue())
        with self.assertLogs("core.views.goz", "ERROR"):
            response = self.client.post(
                reverse("goz_report"), {"archive": empty}, follow=True
            )
        self.assertIn("не найдено ни одного .xls", messages_text(response)[0])
