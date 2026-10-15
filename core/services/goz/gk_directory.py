"""Справочник ГК раздела «Анализ ГОЗ»: поиск записи и загрузка из файлов.

Ключ записи (поле igk) всегда в формате gk_key: «1493», «6766 12»,
«0087/334сб», «0276/481сб 138» и т.п. — см. goz/gk_key.py.

Поддерживаемые файлы справочника:
  * Excel «№ | Предприятие | Заказ | Договор/Дата» — import_extended_rows
    (ГК берётся из договора, заполняются предприятие/заказ/договор);
  * Excel или Word «№ | ГК» — import_simple_values (только ключи ГК).
"""

import re
from decimal import Decimal

import docx
import openpyxl
import xlrd

from ...models import GozContractVat
from .gk_key import gk_key, key_from_contract, parse_gk_key

# Ставка НДС по умолчанию для новых записей справочника
DEFAULT_VAT_RATE = Decimal("22.0")


# --- Чтение справочника ---


def directory_index():
    """Возвращает словарь {ключ ГК: запись} для быстрого поиска.

    Ключи нормализованы: «1234» и «1234/34сб» — разные ГК.
    """
    return {rec.igk: rec for rec in GozContractVat.objects.all()}


# Ключ «ИГК спец» без номера сб: «9741 136»
_BASE_SPEC_RE = re.compile(r"^(\d{4}) (\d+)$")


def spec_counterpart(key, index):
    """Для ключа «9741 136» — запись «9741/NNNсб 136», если она одна.

    В имени файла спецификации номер сб часто не пишется («…09741_136.xls»),
    а в договоре справочника он есть («…09741/255сб/137 от …, спец…/136»).
    Это один и тот же договор. Если подходящих записей несколько — не угадываем.
    """
    m = _BASE_SPEC_RE.match(key)
    if not m:
        return None
    pattern = re.compile(rf"^{m.group(1)}/\d+сб {m.group(2)}$")
    found = [rec for k, rec in index.items() if pattern.match(k)]
    return found[0] if len(found) == 1 else None


def find_gk(label, index):
    """Ищет запись справочника для ГК из имени файла архива.

    1. Точное совпадение ключа.
    2. Для «ИГК спец» без сб — запись «ИГК/NNNсб спец» (см. spec_counterpart).
       Она же выбирается, если точная запись пустая (без договора): такие
       пустые записи раньше создавались при сохранении ставок из архива.

    Args:
        label: метка ГК из файла (например, «9741 136»)
        index: словарь {ключ: запись} от directory_index()

    Returns:
        GozContractVat или None, если ГК в справочнике нет
    """
    key = gk_key(label)
    rec = index.get(key)
    if rec is not None and rec.contract_date:
        return rec
    return spec_counterpart(key, index) or rec


# --- Вспомогательные функции для импорта ---


def _cell(value):
    """Преобразует значение ячейки в строку: None → '', float без дробной части → int."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_header(cells):
    """Проверяет, является ли строка заголовком таблицы (содержит «предприятие» или «заказ»)."""
    low = [c.lower() for c in cells]
    return any("предприятие" in c for c in low) or any(c == "заказ" for c in low)


# --- Импорт расширенного справочника ---


def import_extended_rows(rows):
    """Импортирует расширенный справочник ГК из файла с колонками: № | Предприятие | Заказ | Договор/Дата.

    ГК извлекается из поля «Договор/Дата» (например, «…6066/162сб/137 от …» → «6066/162сб»).
    ГК «6066» и «6066/162сб» считаются разными записями.

    Args:
        rows: итератор строк файла (списки значений ячеек)

    Returns:
        tuple: (добавлено, обновлено, список_ошибок)
    """
    by_key = directory_index()
    added = updated = 0
    errors = []

    for n, row in enumerate(rows, 1):
        # Берём первые 4 колонки, дополняем пустыми при необходимости
        cells = [_cell(c) for c in list(row)[:4]]
        cells += [""] * (4 - len(cells))
        _num, enterprise, order, contract = cells

        # Пропускаем пустые строки и заголовки
        if not any(cells[1:]) or _is_header(cells):
            continue

        # Извлекаем ключ ГК из поля «Договор/Дата»
        key = key_from_contract(contract)
        if not key:
            errors.append(f"строка {n}: в «{contract[:40]}» не найден номер ГК")
            continue

        fields = {
            "enterprise": enterprise[:255] or None,
            "order": order[:100] or None,
            "contract_date": contract[:255] or None,
        }

        rec = by_key.get(key)

        try:
            if rec:
                # Обновляем существующую запись
                for name, value in fields.items():
                    setattr(rec, name, value)
                rec.save()
                updated += 1
            else:
                # Создаём новую запись со ставкой НДС по умолчанию
                by_key[key] = GozContractVat.objects.create(
                    igk=key, vat_rate=DEFAULT_VAT_RATE, **fields
                )
                added += 1
        except Exception as e:
            errors.append(f"строка {n} ({key}): {e}")

    return added, updated, errors


# --- Чтение файлов справочника ---


def read_sheet_rows(file, ext):
    """Строки первого листа Excel (.xlsx или .xls) как кортежи значений."""
    if ext == ".xlsx":
        wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        return [tuple(r) for r in wb.active.iter_rows(min_row=1, values_only=True)]

    ws = xlrd.open_workbook(file_contents=file.read()).sheet_by_index(0)
    return [tuple(ws.row_values(rx)) for rx in range(ws.nrows)]


def is_extended_directory(rows):
    """True, если это файл «№ | Предприятие | Заказ | Договор/Дата».

    Признак — заголовок со словами «предприятие»/«договор» в первых 10 строках
    либо ГК, распознанный в 4-й колонке (файл без заголовка).
    """
    for row in rows[:10]:
        text = " ".join(str(c).lower() for c in row if c is not None)
        if "предприятие" in text or "договор" in text:
            return True
        if len(row) >= 4 and parse_gk_key(row[3]):
            return True
    return False


# Значения второй колонки, которые являются заголовком, а не номером ГК
_HEADER_WORDS = {"", "nan", "none", "гк", "номер"}


def simple_values_from_rows(rows):
    """Номера ГК из Excel «№ | ГК» (вторая колонка).

    Номер, записанный в Excel числом, теряет ведущие нули (0509 -> 509),
    поэтому короткие числа дополняются нулями до 4 цифр.
    """
    values = []
    for row in rows:
        if len(row) < 2 or row[1] is None:
            continue
        value = row[1]
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            value = str(int(value)).zfill(4)
        value = str(value).strip()
        if value.lower() in _HEADER_WORDS:
            continue
        values.append(value)
    return values


def simple_values_from_docx(file):
    """Номера ГК из таблиц Word «№ | ГК» (вторая колонка)."""
    values = []
    for table in docx.Document(file).tables:
        for row in table.rows:
            cells = row.cells
            if len(cells) < 2:
                continue

            num_text = cells[0].text.strip()
            igk = cells[1].text.strip()

            # Строка заголовка («ГК», «Номер ГК») или пустая ячейка
            if not igk or "гк" in igk.lower() or "номер" in igk.lower():
                continue
            # В первой колонке должен быть порядковый номер (или пусто/«№»)
            if num_text and not num_text.isdigit() and num_text.lower() not in ("№", "п/п"):
                continue
            values.append(igk)
    return values


def import_simple_values(values):
    """Добавляет в справочник ГК, которых там ещё нет.

    Существующие записи не меняются (раньше повторная загрузка сбрасывала
    у них ставку НДС на 22%).

    Returns:
        (добавлено, уже_было, ошибки)
    """
    index = directory_index()
    added = existing = 0
    errors = []

    for value in values:
        key = gk_key(value)
        if key in index:
            existing += 1
            continue
        try:
            index[key] = GozContractVat.objects.create(igk=key, vat_rate=DEFAULT_VAT_RATE)
            added += 1
        except Exception as e:
            errors.append(f"{key}: {e}")

    return added, existing, errors
