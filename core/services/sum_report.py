import io
import re
import zipfile
from collections import defaultdict
from copy import copy
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import List

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, PatternFill, Side

TEMPLATE_PATH = (
    Path(__file__).resolve().parents[1] / "static" / "files" / "templateSum.xlsx"
)

COL_IGK = 1
COL_COUNTERAGENT = 2
COL_INN = 3
COL_DEP = 4
COL_CONTRACT = 5
COL_CONDITION = 6
COL_STAGE = 7
COL_OBJECT = 8
COL_CONTRACT_SUM = 10
COL_PLAN_AVANS = 11
COL_PERCENT_DOC = 12
COL_FACT_PAID = 14
COL_PERIOD_REG_DATE = 18
COL_NOTE = 20
COL_PLAN_DATE_CONTRACT = 21
COL_REMAINS_PAY = 24
COL_COMPLETED_SUM = 26
COL_ZNP_COUNT = 27
COL_SUM_80 = 28
COL_PAID_FROM_ZNP = 29
COL_SUM_ISSUED_ZNP = 32
COL_SUM_AVANS = 33

MAX_INPUT_COLUMNS = 50

SMALL_CONTRACT_LIMIT = 1_000_000.0
ALLOWED_DIFFERENCE = 1.0

REPORT_FONT_SIZE = 12
REPORT_END_COLUMN = 11

RED_FILL = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="FF000000"),
    right=Side(style="thin", color="FF000000"),
    top=Side(style="thin", color="FF000000"),
    bottom=Side(style="thin", color="FF000000"),
)


@dataclass
class SumReportRow:
    sheet_name: str = ""
    dep: str = ""
    igk: str = ""
    counteragent: str = ""
    inn: str = ""
    contract: str = ""
    condition: str = ""
    stage: str = ""
    object: str = ""
    order_doc: str = ""
    contract_sum: float = 0.0
    plan_avans: float = 0.0
    percent_doc: float = 0.0
    fact_paid: float = 0.0
    note: str = ""
    completed_sum: float = 0.0
    znp_count: int = 0
    sum_80: float = 0.0
    paid_from_znp: float = 0.0
    remains_pay: float = 0.0
    sum_avans: float = 0.0
    sum_issued_znp: float = 0.0
    period_reg_date: str = ""
    plan_date_contract: str = ""
    is_cycle: bool = False
    igk_title: str = ""


def _to_float(val) -> float:
    if val is None or isinstance(val, bool):
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    try:
        cleaned = str(val).replace(" ", "").replace("\xa0", "").replace(",", ".")
        if cleaned in ("", "-", "None", "FALSE", "False", "false"):
            return 0.0
        return float(cleaned)
    except (ValueError, TypeError):
        return 0.0


def _to_int(val) -> int:
    if val is None or isinstance(val, bool):
        return 0
    if isinstance(val, (int, float)):
        return int(val)
    try:
        cleaned = str(val).replace(" ", "").replace(",", ".")
        if cleaned in ("", "-", "None", "FALSE", "False", "false"):
            return 0
        return int(float(cleaned))
    except (ValueError, TypeError):
        return 0


def _to_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    if isinstance(val, int):
        return str(val)
    return re.sub(r"\s+", " ", str(val)).strip()


def _to_date_str(val) -> str:
    if val is None:
        return ""
    if isinstance(val, datetime):
        if val.year == 1900:
            return ""
        return val.strftime("%d.%m.%Y")
    text = str(val).strip()
    if not text or text.lower() in ("none", "false", "0"):
        return ""
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d.%m.%y", "%m/%d/%Y", "%m/%d/%y"):
        try:
            dt = datetime.strptime(text, fmt)
            if dt.year == 1900:
                return ""
            return dt.strftime("%d.%m.%Y")
        except ValueError:
            continue
    return ""


def _normalize_condition(value: str) -> str:
    if not value:
        return ""
    text = value.replace("\xa0", " ").strip().lower()
    text = re.sub(r"\s+", " ", text)
    text = re.sub(r"\s*[\(\[,].*$", "", text)
    return text.strip()


def _classify_row(row_vals: list) -> str:
    first_three = [str(v).strip().lower() for v in row_vals[:3] if v is not None]
    text = " ".join(first_three)
    if re.search(r"(^|\s)(общий\s+итог|итого\s+по\s+гк|всего)\s*:?\s*$", text):
        return "stop"
    if re.search(r"(^|\s)итого(\s|:)", text):
        return "skip"
    if "в том числе" in text:
        return "skip"
    return "data"


def _is_valid_sheet_name(name: str) -> bool:
    return name.strip().isdigit() and len(name.strip()) == 4


def parse_sum_report_file(file_stream) -> List[SumReportRow]:
    wb = load_workbook(file_stream, data_only=True)
    all_rows: List[SumReportRow] = []

    for ws in wb.worksheets:
        sheet_name = ws.title
        if not _is_valid_sheet_name(sheet_name):
            continue

        igk_title = _to_str(ws["B1"].value)
        header_7 = str(ws.cell(2, COL_STAGE).value or "").lower()
        is_cycle_format = "длинноцикловые" in header_7

        for r in range(3, ws.max_row + 1):
            row_vals = [ws.cell(r, c).value for c in range(1, MAX_INPUT_COLUMNS + 1)]

            row_type = _classify_row(row_vals)
            if row_type == "stop":
                break
            if row_type == "skip":
                continue
            if all(v is None or str(v).strip() == "" for v in row_vals):
                continue

            dep = _to_str(row_vals[COL_DEP - 1])
            if not dep:
                continue

            condition = _to_str(row_vals[COL_CONDITION - 1])
            stage = _to_str(row_vals[COL_STAGE - 1])

            is_cycle = False
            if is_cycle_format:
                is_cycle = stage.lower() in ("да", "yes", "true", "1")

            all_rows.append(
                SumReportRow(
                    sheet_name=sheet_name,
                    dep=dep,
                    igk=_to_str(row_vals[COL_IGK - 1]),
                    counteragent=_to_str(row_vals[COL_COUNTERAGENT - 1]),
                    inn=_to_str(row_vals[COL_INN - 1]),
                    contract=_to_str(row_vals[COL_CONTRACT - 1]),
                    condition=condition,
                    stage=stage,
                    object=_to_str(row_vals[COL_OBJECT - 1]),
                    order_doc="",
                    contract_sum=_to_float(row_vals[COL_CONTRACT_SUM - 1]),
                    plan_avans=_to_float(row_vals[COL_PLAN_AVANS - 1]),
                    percent_doc=_to_float(row_vals[COL_PERCENT_DOC - 1]),
                    fact_paid=_to_float(row_vals[COL_FACT_PAID - 1]),
                    note=_to_str(row_vals[COL_NOTE - 1]),
                    completed_sum=_to_float(row_vals[COL_COMPLETED_SUM - 1]),
                    znp_count=_to_int(row_vals[COL_ZNP_COUNT - 1]),
                    sum_80=_to_float(row_vals[COL_SUM_80 - 1]),
                    paid_from_znp=_to_float(row_vals[COL_PAID_FROM_ZNP - 1]),
                    remains_pay=_to_float(row_vals[COL_REMAINS_PAY - 1]),
                    sum_avans=_to_float(row_vals[COL_SUM_AVANS - 1]),
                    sum_issued_znp=_to_float(row_vals[COL_SUM_ISSUED_ZNP - 1]),
                    period_reg_date=_to_date_str(row_vals[COL_PERIOD_REG_DATE - 1]),
                    plan_date_contract=_to_date_str(
                        row_vals[COL_PLAN_DATE_CONTRACT - 1]
                    ),
                    is_cycle=is_cycle,
                    igk_title=igk_title,
                )
            )

    return all_rows


def _safe_sheet_title(value):
    invalid_chars = set("[]:*?/\\")
    result = "".join(ch for ch in str(value) if ch not in invalid_chars).strip()
    return result[:31] or "Report"


def _extract_year_from_title(title):
    match = re.search(r"20\d{2}", title)
    return match.group(0) if match else str(timezone.localdate().year)


def _find_row_by_marker(ws, marker, start_row=1, end_row=None, columns=(1, 2, 3)):
    end = end_row or ws.max_row
    marker_l = marker.lower()
    for r in range(start_row, end + 1):
        for c in columns:
            val = ws.cell(r, c).value
            if val and marker_l in str(val).lower():
                return r
    return None


def _delete_rows(ws, start_row, end_row):
    if start_row > end_row:
        return
    ws.delete_rows(start_row, end_row - start_row + 1)


def _prepare_data_area(ws, start_row=5):
    to_unmerge = [str(mr) for mr in ws.merged_cells.ranges if mr.min_row >= start_row]
    for rng in to_unmerge:
        ws.unmerge_cells(rng)

    for r in list(ws.row_dimensions.keys()):
        if r >= start_row:
            del ws.row_dimensions[r]


_HEADER_OR_TOTAL_KEYWORDS = (
    "итого",
    "всего",
    "в том числе",
    "общий итог",
    "кол-во",
    "отдел",
    "контрагент",
    "предмет",
    "краткая справка",
    "наименование",
    "показатель",
)


def _is_header_or_total_row(ws, r) -> bool:
    for c in (1, 2, 3):
        val = ws.cell(row=r, column=c).value
        if val is None:
            continue
        text = str(val).strip().lower()
        for kw in _HEADER_OR_TOTAL_KEYWORDS:
            if kw in text:
                return True
    return False


def _find_first_data_row_table1(ws, total_row_idx):
    first_data_row = total_row_idx
    r = total_row_idx - 1
    while r > 1:
        if _is_header_or_total_row(ws, r):
            break
        if any(ws.cell(row=r, column=c).value is not None for c in (1, 2, 3)):
            first_data_row = r
        r -= 1
    return first_data_row


def _find_first_data_row_detail(ws, total_row_idx):
    first_data_row = total_row_idx
    r = total_row_idx - 1
    while r > 1:
        if _is_header_or_total_row(ws, r):
            break
        if any(ws.cell(row=r, column=c).value is not None for c in (1, 2)):
            first_data_row = r
        r -= 1
    return first_data_row


def _capture_styles(ws, row_idx, columns):
    styles = []
    for c in columns:
        cell = ws.cell(row=row_idx, column=c)
        styles.append(
            {
                "font": copy(cell.font),
                "border": copy(cell.border),
                "fill": copy(cell.fill),
                "number_format": cell.number_format,
                "protection": copy(cell.protection),
                "alignment": copy(cell.alignment),
            }
        )
    return styles


def _apply_style(cell, style):
    cell.font = copy(style["font"])
    cell.border = copy(style["border"])
    cell.fill = copy(style["fill"])
    cell.number_format = style["number_format"]
    cell.protection = copy(style["protection"])
    cell.alignment = copy(style["alignment"])


def _finalize_sheet(
    ws, start_row=3, end_col=REPORT_END_COLUMN, font_size=REPORT_FONT_SIZE
):
    end_row = ws.max_row
    for r in range(start_row, end_row + 1):
        has_content = False
        for c in range(1, end_col + 1):
            v = ws.cell(row=r, column=c).value
            if v is not None and str(v).strip() != "":
                has_content = True
                break
        if not has_content:
            continue
        for c in range(1, end_col + 1):
            cell = ws.cell(row=r, column=c)
            f = copy(cell.font)
            f.size = font_size
            cell.font = f
            cell.border = THIN_BORDER


def _aggregate_table1(rows):
    deps = defaultdict(
        lambda: {
            "contract_count": 0,
            "contract_sum": 0.0,
            "completed_count": 0,
            "completed_sum": 0.0,
            "plan_avans": 0.0,
            "znp_count": 0,
            "znp_sum": 0.0,
            "znp_paid_sum": 0.0,
            "remain_pay": 0.0,
            "total_remain_pay": 0.0,
        }
    )

    for r in rows:
        d = deps[r.dep]
        d["contract_count"] += 1
        d["contract_sum"] += r.contract_sum

        if _normalize_condition(r.condition) == "заключен":
            d["completed_count"] += 1
            d["completed_sum"] += r.contract_sum

        d["plan_avans"] += r.sum_80

        if r.znp_count != 0:
            d["znp_count"] += 1
            d["znp_sum"] += r.completed_sum
            d["znp_paid_sum"] += r.paid_from_znp
            if r.remains_pay != 0:
                d["remain_pay"] += abs(r.sum_avans)

        d["total_remain_pay"] += r.remains_pay

    return deps


def _write_table1(ws, deps, total_row_idx):
    if not deps:
        return total_row_idx

    first_data_row = _find_first_data_row_table1(ws, total_row_idx)
    style_row = total_row_idx - 1 if first_data_row < total_row_idx else total_row_idx
    data_styles = _capture_styles(ws, style_row, range(1, 12))

    if first_data_row < total_row_idx:
        _delete_rows(ws, first_data_row, total_row_idx - 1)
        total_row_idx = first_data_row

    sorted_deps = sorted(deps.keys(), key=lambda x: (len(x), x))
    n = len(sorted_deps)

    if n > 0:
        ws.insert_rows(total_row_idx, amount=n)

    current_row = total_row_idx
    for dep in sorted_deps:
        data = deps[dep]
        values = [
            dep,
            data["contract_count"],
            round(data["contract_sum"], 2),
            data["completed_count"],
            round(data["completed_sum"], 2),
            round(data["plan_avans"], 2),
            data["znp_count"],
            round(data["znp_sum"], 2),
            round(data["znp_paid_sum"], 2),
            round(data["remain_pay"], 2),
            round(data["total_remain_pay"], 2),
        ]
        for idx, value in enumerate(values, 1):
            cell = ws.cell(row=current_row, column=idx)
            _apply_style(cell, data_styles[idx - 1])
            cell.value = value
        current_row += 1

    total_row = current_row

    totals = {
        "contract_count": sum(d["contract_count"] for d in deps.values()),
        "contract_sum": round(sum(d["contract_sum"] for d in deps.values()), 2),
        "completed_count": sum(d["completed_count"] for d in deps.values()),
        "completed_sum": round(sum(d["completed_sum"] for d in deps.values()), 2),
        "plan_avans": round(sum(d["plan_avans"] for d in deps.values()), 2),
        "znp_count": sum(d["znp_count"] for d in deps.values()),
        "znp_sum": round(sum(d["znp_sum"] for d in deps.values()), 2),
        "znp_paid_sum": round(sum(d["znp_paid_sum"] for d in deps.values()), 2),
        "remain_pay": round(sum(d["remain_pay"] for d in deps.values()), 2),
        "total_remain_pay": round(sum(d["total_remain_pay"] for d in deps.values()), 2),
    }

    for idx, key in enumerate(
        [
            "contract_count",
            "contract_sum",
            "completed_count",
            "completed_sum",
            "plan_avans",
            "znp_count",
            "znp_sum",
            "znp_paid_sum",
            "remain_pay",
            "total_remain_pay",
        ],
        2,
    ):
        ws.cell(row=total_row, column=idx).value = totals[key]

    return total_row


def _filter_table2(rows):
    result = []
    for r in rows:
        status = _normalize_condition(r.condition)
        if status in ("знп оформлена", "заключен"):
            if not (-1 <= r.remains_pay <= 1) and not (-1 <= r.sum_avans <= 1):
                result.append(r)
    return result


def _filter_table3(rows):
    result = []
    for r in rows:
        status = _normalize_condition(r.condition)
        if status in ("знп на оформлении", "заключен"):
            if not (-1 <= r.remains_pay <= 1) and not (
                -1 <= round(r.sum_issued_znp, 2) <= 1
            ):
                result.append(r)
    return result


def _filter_table4(rows):
    result = []
    for r in rows:
        if _normalize_condition(r.condition) == "не заключен":
            result.append(r)
    return result


def _write_detail_table(ws, total_row_idx, rows, table_kind="table2"):
    if not rows:
        return total_row_idx, 0.0, 0.0

    if table_kind == "table4":

        def avans_getter(r):
            return round(abs(r.remains_pay), 2)

        def sixth_getter(r):
            return r.plan_date_contract

    elif table_kind == "table3":

        def avans_getter(r):
            return round(abs(r.sum_issued_znp), 2)

        def sixth_getter(r):
            return r.condition

    else:

        def avans_getter(r):
            return round(abs(r.sum_avans), 2)

        def sixth_getter(r):
            return r.condition

    large_rows = []
    other_contract_sum = 0.0
    other_avans_sum = 0.0

    for r in rows:
        contract_sum = round(r.contract_sum, 2)
        sum_avans = avans_getter(r)
        if contract_sum < SMALL_CONTRACT_LIMIT:
            other_contract_sum += contract_sum
            other_avans_sum += sum_avans
        else:
            large_rows.append((r, contract_sum, sum_avans))

    first_data_row = _find_first_data_row_detail(ws, total_row_idx)
    style_row = total_row_idx - 1 if first_data_row < total_row_idx else total_row_idx
    data_styles = _capture_styles(ws, style_row, range(1, 10))

    if first_data_row < total_row_idx:
        _delete_rows(ws, first_data_row, total_row_idx - 1)
        total_row_idx = first_data_row

    has_others = other_contract_sum != 0 or other_avans_sum != 0
    total_new_rows = len(large_rows) + (1 if has_others else 0)

    if total_new_rows > 0:
        ws.insert_rows(total_row_idx, amount=total_new_rows)

    current_row = total_row_idx
    for r, contract_sum, sum_avans in large_rows:
        percent_val = round(sum_avans / contract_sum * 100, 0) if contract_sum else 0
        values = [
            r.counteragent,
            r.object,
            contract_sum,
            sum_avans,
            f"{int(percent_val)}%",
            sixth_getter(r),
            r.period_reg_date,
            r.note,
            r.dep,
        ]
        for idx, value in enumerate(values, 1):
            cell = ws.cell(row=current_row, column=idx)
            style = data_styles[idx - 1]
            cell.font = copy(style["font"])
            cell.border = copy(style["border"])
            cell.fill = copy(style["fill"])
            cell.number_format = style["number_format"]
            cell.protection = copy(style["protection"])
            if idx in (1, 2, 6, 8, 9):
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            else:
                cell.alignment = copy(style["alignment"])
            cell.value = value
        current_row += 1

    if has_others:
        for idx in range(1, 10):
            cell = ws.cell(row=current_row, column=idx)
            _apply_style(cell, data_styles[idx - 1])
        ws.cell(row=current_row, column=1).value = "Прочие контрагенты:"
        ws.cell(row=current_row, column=3).value = round(other_contract_sum, 2)
        ws.cell(row=current_row, column=4).value = round(other_avans_sum, 2)
        current_row += 1

    total_row = current_row
    total_contract = round(sum(c for _, c, _ in large_rows) + other_contract_sum, 2)
    total_avans = round(sum(a for _, _, a in large_rows) + other_avans_sum, 2)

    ws.cell(row=total_row, column=3).value = total_contract
    ws.cell(row=total_row, column=4).value = total_avans

    return total_row, total_contract, total_avans


def _adjust_column_widths(ws):
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 20
    ws.column_dimensions["G"].width = 15
    ws.column_dimensions["H"].width = 40
    ws.column_dimensions["I"].width = 10


def generate_single_report(rows, igk_title, is_cycle=False):
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Шаблон {TEMPLATE_PATH} не найден")

    wb = load_workbook(TEMPLATE_PATH)
    ws = wb["9022"] if "9022" in wb.sheetnames else wb.active

    _prepare_data_area(ws, start_row=5)

    year = _extract_year_from_title(igk_title)
    ws["A2"] = igk_title.replace("__g", "").replace("__y", year).strip()
    ws.cell(row=1, column=11).value = timezone.localdate().strftime("%d.%m.%Y")

    t1_total = _find_row_by_marker(ws, "Всего", start_row=5, end_row=10) or 6
    deps = _aggregate_table1(rows)
    t1_total = _write_table1(ws, deps, t1_total)

    t2_total = _find_row_by_marker(
        ws, "ИТОГО:", start_row=t1_total + 1, end_row=t1_total + 15
    ) or (t1_total + 4)
    rows2 = _filter_table2(rows)
    t2_total, contract_sum_total, avans_sum_total = _write_detail_table(
        ws, t2_total, rows2, table_kind="table2"
    )

    t3_total = _find_row_by_marker(
        ws, "ИТОГО:", start_row=t2_total + 1, end_row=t2_total + 15
    ) or (t2_total + 3)
    rows3 = _filter_table3(rows)
    t3_total, c3, a3 = _write_detail_table(ws, t3_total, rows3, table_kind="table3")
    contract_sum_total = round(contract_sum_total + c3, 2)
    avans_sum_total = round(avans_sum_total + a3, 2)

    t4_total = _find_row_by_marker(
        ws, "ИТОГО:", start_row=t3_total + 1, end_row=t3_total + 15
    ) or (t3_total + 3)
    rows4 = _filter_table4(rows)
    t4_total, c4, a4 = _write_detail_table(ws, t4_total, rows4, table_kind="table4")
    contract_sum_total = round(contract_sum_total + c4, 2)
    avans_sum_total = round(avans_sum_total + a4, 2)

    final_total_row = t4_total + 1
    ws.cell(row=final_total_row, column=3).value = contract_sum_total
    ws.cell(row=final_total_row, column=4).value = avans_sum_total

    final_title_cell = ws.cell(row=final_total_row, column=1)
    if isinstance(final_title_cell.value, str):
        final_title_cell.value = final_title_cell.value.replace("__g", igk_title)

    val10 = round(float(ws.cell(row=t1_total, column=10).value or 0), 2)
    val_sec4 = round(float(ws.cell(row=t2_total, column=4).value or 0), 2)
    if abs(val10 - val_sec4) > ALLOWED_DIFFERENCE:
        ws.cell(row=t1_total, column=10).fill = RED_FILL

    val11 = round(float(ws.cell(row=t1_total, column=11).value or 0), 2)
    val_last4 = round(float(ws.cell(row=final_total_row, column=4).value or 0), 2)
    if abs(val11 - val_last4) > ALLOWED_DIFFERENCE:
        ws.cell(row=t1_total, column=11).fill = RED_FILL

    _finalize_sheet(ws, start_row=3, end_col=REPORT_END_COLUMN)
    _adjust_column_widths(ws)
    ws.title = _safe_sheet_title(igk_title[:31])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_sum_reports_zip(file_stream, is_cycle=False, username="user"):
    all_rows = parse_sum_report_file(file_stream)
    if not all_rows:
        raise ValueError("Нет данных для формирования отчёта.")

    sheet_groups = defaultdict(list)
    sheet_titles = {}
    for row in all_rows:
        sheet_groups[row.sheet_name].append(row)
        sheet_titles[row.sheet_name] = row.igk_title

    report_label = "Краткая (Длинноцикл.)" if is_cycle else "Краткая"
    today = timezone.localdate().strftime("%d-%m-%Y")
    buffer = io.BytesIO()

    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for sheet_name, rows in sheet_groups.items():
            try:
                data = generate_single_report(rows, sheet_titles[sheet_name], is_cycle)
            except Exception as exc:
                raise RuntimeError(
                    f"Ошибка формирования отчёта ИГК {sheet_name}: {exc}"
                ) from exc

            filename = f"{sheet_name} {report_label} {today}_{username}.xlsx"
            archive.writestr(filename, data)

    return buffer.getvalue()


def generate_sum_reports_response(file_stream, is_cycle=False, username="user"):
    data = build_sum_reports_zip(file_stream, is_cycle, username)
    today = timezone.localdate().strftime("%d-%m-%Y")
    suffix = "cycle" if is_cycle else "normal"

    response = HttpResponse(data, content_type="application/zip")
    response["Content-Disposition"] = (
        f'attachment; filename="reports_{suffix}_{today}.zip"'
    )
    return response
