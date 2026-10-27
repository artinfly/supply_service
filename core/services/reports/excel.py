"""Создание Excel-файлов с форматированием и формирование HTTP-ответов для скачивания."""

import re
from io import BytesIO
from urllib.parse import quote

from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Стили ячеек ---


_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

_HEADER_FONT = Font(name="Arial", size=9, bold=True, color="000000")
_HEADER_FILL = PatternFill("solid", fgColor="FFFFFF")

_NORMAL_FONT = Font(name="Arial", size=9)
_BOLD_FONT = Font(name="Arial", size=9, bold=True, color="000000")
_TOTAL_FILL = PatternFill("solid", fgColor="C0C0C0")

_ALIGN_HEADER = Alignment(horizontal="center", wrap_text=True)
_ALIGN_LEFT = Alignment(horizontal="left")
_ALIGN_RIGHT = Alignment(horizontal="right")

# Таблица транслитерации для ASCII-имён файлов (кириллица → латиница)
_TRANSLIT = str.maketrans(
    "абвгдеёжзийклмнопрстуфхцчшщъыьэюяАБВГДЕЁЖЗИЙКЛМНОПРСТУФХЦЧШЩЪЫЬЭЮЯ ",
    "abvgdeejzijklmnoprstufxccssiieeuaABVGDEEJZIJKLMNOPRSTUFXCCSSIIEEUA_",
)


# --- Вспомогательные функции ---


def _get_number_format(value):
    """Возвращает формат числа: '0' для целых, '0.00' для дробных."""
    if isinstance(value, int) or (isinstance(value, float) and value.is_integer()):
        return "0"
    return "0.00"


# --- Создание Excel-файлов ---


def make_wb(sheet_name, headers, col_widths, rows_data, kinds=None, formats=None):
    """Создаёт Workbook с заголовками, данными и форматированием.

    Args:
        sheet_name: имя листа
        headers: список заголовков колонок
        col_widths: список ширин колонок
        rows_data: список строк данных (списки значений)
        kinds: список типов строк ('normal', 'subtotal', 'total') для форматирования
        formats: словарь {индекс_колонки: формат_числа} для переопределения форматов

    Returns:
        bytes: содержимое Excel-файла
    """
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.freeze_panes = "A2"  # Фиксируем первую строку (заголовки)

    # Форматируем заголовки
    for ci, (header, width) in enumerate(zip(headers, col_widths, strict=True), 1):
        cell = ws.cell(row=1, column=ci, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.border = _BORDER
        cell.alignment = _ALIGN_HEADER
        ws.column_dimensions[get_column_letter(ci)].width = width

    # Форматируем данные
    for ri, row in enumerate(rows_data, 2):
        kind = kinds[ri - 2] if kinds else "normal"

        for ci, value in enumerate(row, 1):
            cell = ws.cell(row=ri, column=ci, value=value)
            cell.border = _BORDER

            is_num = isinstance(value, (int, float)) and not isinstance(value, bool)

            # Форматирование для итоговых строк
            if kind in ("total", "subtotal"):
                cell.font = _BOLD_FONT
                cell.fill = _TOTAL_FILL
                cell.alignment = _ALIGN_RIGHT if is_num else _ALIGN_LEFT
            else:
                cell.font = _NORMAL_FONT
                if is_num:
                    cell.alignment = _ALIGN_RIGHT

            # Формат чисел
            if is_num:
                if formats and ci in formats:
                    cell.number_format = formats[ci]
                else:
                    cell.number_format = _get_number_format(value)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- HTTP-ответы для скачивания ---


def xlsx_response(data, filename_ru):
    """Формирует HttpResponse для скачивания Excel-файла.

    Args:
        data: bytes содержимого Excel-файла
        filename_ru: имя файла на русском (будет транслитерировано для ASCII)

    Returns:
        HttpResponse с правильными заголовками Content-Disposition
    """
    today = timezone.localdate().strftime("%d_%m_%Y")

    # Символы, недопустимые в имени файла Windows и в заголовке ответа
    # (в имени бывает название контрагента: ООО "Ромашка")
    filename_ru = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "_", filename_ru)

    # ASCII-версия имени для старых клиентов: всё, что не стало латиницей
    # после транслитерации («», № и т.п.), заменяем на «_»
    fname_ascii = re.sub(
        r"[^A-Za-z0-9._-]", "_", f"{filename_ru.translate(_TRANSLIT)}_{today}.xlsx"
    )
    # UTF-8 версия для современных клиентов (RFC 5987)
    fname_utf8 = quote(f"{filename_ru}_{today}.xlsx")

    response = HttpResponse(
        data,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    # Двойное имя файла: ASCII fallback + UTF-8 для корректного отображения кириллицы
    response["Content-Disposition"] = (
        f"attachment; filename=\"{fname_ascii}\"; filename*=UTF-8''{fname_utf8}"
    )
    return response
