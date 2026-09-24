"""
Страницы приложения: сводки, реестры, загрузка файлов.

Каждая страница отдаёт только каркас (шаблон). Данные для таблиц
подгружаются через /reports/api/... и рисуются на клиенте.
Сводки (плашки и таблицы по ЦФО) считаются на сервере.
"""

import json
import os
import re
import tempfile
import zipfile
from datetime import datetime
from decimal import Decimal
from functools import wraps
from io import StringIO

import docx
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.core.management import call_command
from django.db import connection
from django.db.models import Count, Exists, OuterRef, Q, Sum
from django.db.models.expressions import RawSQL
from django.shortcuts import redirect, render
from django.utils import timezone

from ..models import (
    GozContractVat,
    IgkStatData,
    NsiIgk,
    SystemEvent,
    ZnpData,
    ZnpDataSAP,
)
from ..services import goz_analysis
from ..services.dashboards import (
    EMPTY_CFO_STATS,
    EMPTY_NOT_ISSUED,
    EMPTY_STAGES,
    EMPTY_ZNP,
    ZNP_STAGE_LABELS,
    ZNP_STAGE_NAMES,
    breakdown_from_stats,
    cfo_breakdown_row,
    cfo_row,
    cfo_totals_row,
    filter_by_year,
    not_issued_aggregates,
    percent,
    sap_cards,
    stage_aggregates,
    to_mln,
    znp_aggregates,
)
from ..services.excel import xlsx_response
from ..services.excel_import import CONTRACT_COLUMNS, ZNP_COLUMNS, ZNP_SAP_COLUMNS
from ..services.queries import (
    ADVANCE,
    CONCLUDED,
    NOT_CONCL,
    SAP_CFO,
    TERMINATED,
    YEARS,
    distinct_agents,
    distinct_cfo,
    distinct_igk_suffixes,
    distinct_sap_cfo,
    distinct_sap_igk,
    needs_znp,
    valid_date,
    valid_year,
)
from ..services.sap_status import (
    SAP_STAGE_NAMES,
    SAP_STAGE_PARAMS,
    sap_second_date,
    sap_status_expr,
)

# --- Общие вспомогательные функции ---


def _ctx(request):
    """Базовый контекст для всех шаблонов: список годов и колонок-флагов."""
    return {
        "years": YEARS,
        "year_cols": [(y, f"y{str(y)[2:]}") for y in YEARS],
    }


def superuser_required(view_func):
    """Пускает только суперпользователей, остальных возвращает на «Анализ ГОЗ»."""

    @wraps(view_func)
    def wrapper(request, *args, **kwargs):
        if not request.user.is_superuser:
            messages.error(
                request, "Загрузка справочника доступна только суперпользователям."
            )
            return redirect("goz_report")
        return view_func(request, *args, **kwargs)

    return wrapper


# Условие «у строки договора есть заказ» (для ORM).
# ВАЖНО: синхронизировать с HAS_ORDER в queries.py
HAS_ORDER_Q = Q(order__isnull=False) & ~Q(order__regex=r"^\s*$")


def normalize_igk(s):
    """Приводит имя ГК к общему виду (убирает слеши, подчёркивания и т.д.)."""
    return re.sub(r"[^0-9a-zA-Z]", "", str(s))


def _igk_and_cfo_lists():
    """Возвращает списки уникальных ИГК и ЦФО для фильтров."""
    with connection.cursor() as cur:
        cur.execute(distinct_igk_suffixes())
        igk_list = [r[0] for r in cur.fetchall()]
        cur.execute(distinct_cfo())
        cfo_list = [r[0] for r in cur.fetchall()]
    return igk_list, cfo_list


# --- Аутентификация ---


def login_view(request):
    """Страница входа."""
    if request.user.is_authenticated:
        return redirect("/reports/")
    error = False
    if request.method == "POST":
        user = authenticate(
            request,
            username=request.POST.get("username"),
            password=request.POST.get("password"),
        )
        if user:
            login(request, user)
            return redirect("/reports/")
        error = True
    return render(request, "login.html", {"error": error})


def logout_view(request):
    """Выход из системы."""
    logout(request)
    return redirect("login")


# --- Главная страница и реестры по годам ---


@login_required
def index(request):
    """Главная страница — меню разделов."""
    return render(request, "index.html", _ctx(request))


@login_required
def kdr_table(request, year):
    """Контроль договорной работы за год."""
    year = valid_year(year)
    ctx = _ctx(request)
    ctx["year"] = year
    return render(request, "kdr_table.html", ctx)


@login_required
def igk_concluded_table(request, year):
    """Заключённые договоры по ИГК за год."""
    year = valid_year(year)
    ctx = _ctx(request)
    ctx.update(
        {"year": year, "report_type": "concluded", "title": f"ИГК {year} — Заключённые"}
    )
    return render(request, "igk_table.html", ctx)


@login_required
def igk_not_concluded_table(request, year):
    """Незаключённые договоры по ИГК за год."""
    year = valid_year(year)
    ctx = _ctx(request)
    ctx.update(
        {
            "year": year,
            "report_type": "not_concluded",
            "title": f"ИГК {year} — Незаключённые",
        }
    )
    return render(request, "igk_table.html", ctx)


@login_required
def igk_terminated_table(request, year):
    """Расторгнутые договоры по ИГК за год."""
    year = valid_year(year)
    ctx = _ctx(request)
    ctx.update(
        {
            "year": year,
            "report_type": "terminated",
            "title": f"ИГК {year} — Расторгнутые",
        }
    )
    return render(request, "igk_table.html", ctx)


# --- Реестры: каркасы страниц ---


@login_required
def all_contracts_table(request):
    """Реестр всех договоров с фильтрами."""
    igk_list, cfo_list = _igk_and_cfo_lists()
    ctx = _ctx(request)
    ctx.update(
        {
            "igk_list": igk_list,
            "cfo_list": cfo_list,
            "concluded_statuses": list(CONCLUDED),
            "not_concl_statuses": list(NOT_CONCL),
            "terminated_statuses": list(TERMINATED),
        }
    )
    return render(request, "all_contracts.html", ctx)


@login_required
def znp_list_table(request):
    """Реестр заявок ФЗД."""
    igk_list, cfo_list = _igk_and_cfo_lists()
    ctx = _ctx(request)
    ctx.update({"igk_list": igk_list, "cfo_list": cfo_list})
    return render(request, "znp_list.html", ctx)


@login_required
def history_status_table(request):
    """История изменений статуса договора."""
    return render(request, "history_status.html", _ctx(request))


@login_required
def history_plan_table(request):
    """История изменений плана."""
    return render(request, "history_plan.html", _ctx(request))


@login_required
def history_fact_table(request):
    """История изменений факта."""
    return render(request, "history_fact.html", _ctx(request))


@login_required
def contract_dupes_table(request):
    """Дубликаты договоров."""
    with connection.cursor() as cur:
        cur.execute(distinct_cfo())
        cfo_list = [r[0] for r in cur.fetchall()]
    ctx = _ctx(request)
    ctx["cfo_list"] = cfo_list
    return render(request, "contract_dupes.html", ctx)


@login_required
def export_page(request):
    """Страница со списком доступных Excel-выгрузок."""
    with connection.cursor() as cur:
        cur.execute(distinct_agents())
        agents = [r[0] for r in cur.fetchall()]
    ctx = _ctx(request)
    ctx["agents"] = agents
    return render(request, "export.html", ctx)


@login_required
def znp_sap_list_table(request):
    """Реестр заявок SAP."""
    with connection.cursor() as cur:
        cur.execute(distinct_sap_igk())
        igk_list = [r[0] for r in cur.fetchall()]
        cur.execute(distinct_sap_cfo())
        cfo_list = [r[0] for r in cur.fetchall()]
    ctx = _ctx(request)
    ctx.update({"igk_list": igk_list, "cfo_list": cfo_list})
    return render(request, "znp_sap_list.html", ctx)


# --- Загрузка файлов ---

FILE_TYPE_COMMANDS = {
    "contracts": "load_contracts",
    "znp": "load_znp",
    "znp_sap": "load_znp_sap",
}

FILE_TYPE_COLUMNS = {
    "contracts": list(CONTRACT_COLUMNS),
    "znp": list(ZNP_COLUMNS),
    "znp_sap": list(ZNP_SAP_COLUMNS),
}

FILE_TYPE_LABELS = {
    "contracts": "Договоры",
    "znp": "ЗНП (ФЗД)",
    "znp_sap": "ЗНП (SAP)",
}


@login_required
def upload_excel(request):
    """Страница загрузки файлов и обработчик загрузки."""
    result = None
    file_type = request.POST.get("file_type", "contracts")
    if request.method == "POST" and request.FILES.get("excel_file"):
        command = FILE_TYPE_COMMANDS.get(file_type)
        f = request.FILES["excel_file"]
        ext = os.path.splitext(f.name)[1].lower()

        if ext not in (".xlsx", ".xls"):
            messages.error(
                request,
                f"Неподдерживаемый формат файла: {ext or 'без расширения'}. Нужен .xlsx или .xls",
            )
            ctx = _ctx(request)
            ctx["file_types"] = FILE_TYPE_LABELS
            ctx["file_columns"] = FILE_TYPE_COLUMNS
            ctx["selected_type"] = file_type
            return render(request, "upload.html", ctx)

        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
            for chunk in f.chunks():
                tmp.write(chunk)
            tmp_path = tmp.name
        try:
            if command is None:
                raise ValueError(f"Неизвестный тип файла: {file_type}")
            out = StringIO()
            call_command(command, tmp_path, stdout=out)
            result = out.getvalue()
            messages.success(request, "Файл успешно загружен и нормализован")
        except Exception as e:
            text = str(e)
            messages.error(
                request, text if text.startswith("Ошибка") else f"Ошибка: {text}"
            )
            result = str(e)
        finally:
            os.unlink(tmp_path)
    ctx = _ctx(request)
    ctx["result"] = result
    ctx["file_types"] = FILE_TYPE_LABELS
    ctx["file_columns"] = FILE_TYPE_COLUMNS
    ctx["selected_type"] = file_type
    return render(request, "upload.html", ctx)


@login_required
def goz_report(request):
    """Анализ отчётов ЕИС ГОЗ с индивидуальными ставками НДС."""
    ctx = _ctx(request)

    # Список справочника для отображения на странице
    directory_list = GozContractVat.objects.all().order_by("igk", "year")
    ctx["directory_list"] = directory_list

    if request.method == "POST":
        temp_zip_path = request.POST.get("temp_zip_path")
        vat_rates_json = request.POST.get("vat_rates")

        # Шаг 2: Генерация отчёта после настройки ставок
        if vat_rates_json and temp_zip_path and os.path.exists(temp_zip_path):
            try:
                vat_rates = json.loads(vat_rates_json)
                save_to_dir = request.POST.get("save_to_dir") == "on"

                if save_to_dir:
                    for igk, rate in vat_rates.items():
                        try:
                            rate_val = Decimal(str(rate).replace(",", "."))
                            # При сохранении из архива обновляем/создаём запись без года
                            GozContractVat.objects.update_or_create(
                                igk=igk,
                                year=None,
                                defaults={"vat_rate": rate_val},
                            )
                        except Exception:
                            pass

                contracts = goz_analysis.read_archive(temp_zip_path)
                all_vats = list(GozContractVat.objects.all())
                vats_mapping = {}
                for v in all_vats:
                    key = normalize_igk(v.igk)
                    # Берём запись с максимальным годом (если год указан), иначе первую попавшуюся
                    if key not in vats_mapping:
                        vats_mapping[key] = v
                    else:
                        current_year = vats_mapping[key].year or ""
                        new_year = v.year or ""
                        if new_year > current_year:
                            vats_mapping[key] = v

                normalized_contracts = []
                for zip_igk, plan, fact in contracts:
                    norm = normalize_igk(zip_igk)
                    db_vat = vats_mapping.get(norm)
                    db_igk = db_vat.igk if db_vat else zip_igk
                    normalized_contracts.append((db_igk, plan, fact))

                data = goz_analysis.build_report(normalized_contracts, vat_rates)
                os.unlink(temp_zip_path)
                return xlsx_response(data, "Анализ_ГОЗ")
            except Exception as e:
                messages.error(request, f"Ошибка формирования отчёта: {e}")
                if os.path.exists(temp_zip_path):
                    os.unlink(temp_zip_path)

        # Шаг 1: Загрузка ZIP-архива
        elif request.FILES.get("archive"):
            try:
                archive = request.FILES["archive"]
                with tempfile.NamedTemporaryFile(delete=False, suffix=".zip") as tmp:
                    for chunk in archive.chunks():
                        tmp.write(chunk)
                    tmp_path = tmp.name

                contracts = goz_analysis.read_archive(tmp_path)
                gks_in_archive = [c[0] for c in contracts]

                all_vats = list(GozContractVat.objects.all())
                vats_mapping = {}
                for v in all_vats:
                    key = normalize_igk(v.igk)
                    if key not in vats_mapping:
                        vats_mapping[key] = v
                    else:
                        current_year = vats_mapping[key].year or ""
                        new_year = v.year or ""
                        if new_year > current_year:
                            vats_mapping[key] = v

                gk_list = []
                for zip_igk in gks_in_archive:
                    norm = normalize_igk(zip_igk)
                    db_vat = vats_mapping.get(norm)

                    if db_vat:
                        gk_list.append(
                            {
                                "igk": db_vat.igk,
                                "year": db_vat.year or "",
                                "product": db_vat.product or "",
                                "vat_rate": float(db_vat.vat_rate),
                                "is_new": False,
                            }
                        )
                    else:
                        gk_list.append(
                            {
                                "igk": zip_igk,
                                "year": "",
                                "product": "",
                                "vat_rate": 22.0,
                                "is_new": True,
                            }
                        )

                ctx["temp_zip_path"] = tmp_path
                ctx["gk_list"] = gk_list
                ctx["step2"] = True
            except zipfile.BadZipFile:
                messages.error(request, "Файл не является ZIP-архивом")
                if "tmp_path" in locals() and os.path.exists(tmp_path):
                    os.unlink(tmp_path)
            except Exception as e:
                messages.error(request, f"Ошибка обработки архива: {e}")
                if "tmp_path" in locals() and os.path.exists(tmp_path):
                    os.unlink(tmp_path)

    if "step2" not in ctx:
        ctx["default_vat_rate"] = 22.0

    return render(request, "goz_report.html", ctx)


@login_required
@superuser_required
def upload_gk_directory(request):
    """Загрузка справочника ГК из Word или Excel файла."""
    if request.method == "POST" and request.FILES.get("doc_file"):
        file = request.FILES["doc_file"]
        ext = os.path.splitext(file.name)[1].lower()

        added_count = 0
        updated_count = 0

        try:
            if ext == ".docx":
                doc = docx.Document(file)
                for table in doc.tables:
                    for row in table.rows:
                        cells = row.cells
                        if len(cells) >= 2:
                            num_text = cells[0].text.strip()
                            igk = cells[1].text.strip()

                            if not igk or "гк" in igk.lower() or "номер" in igk.lower():
                                continue
                            if (
                                num_text
                                and not num_text.isdigit()
                                and num_text.lower() not in ["№", "п/п"]
                            ):
                                continue

                            try:
                                rate_val = Decimal("22.0")
                                obj, created = GozContractVat.objects.update_or_create(
                                    igk=igk, defaults={"vat_rate": rate_val}
                                )
                                if created:
                                    added_count += 1
                                else:
                                    updated_count += 1
                            except Exception:
                                continue
            elif ext == ".xlsx":
                import openpyxl

                wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
                ws = wb.active
                for row in ws.iter_rows(min_row=1, values_only=True):
                    if len(row) >= 2:
                        igk = str(row[1]).strip() if row[1] is not None else ""
                        if not igk or igk.lower() in ["nan", "гк", "номер", "none"]:
                            continue
                        try:
                            rate_val = Decimal("22.0")
                            obj, created = GozContractVat.objects.update_or_create(
                                igk=igk, defaults={"vat_rate": rate_val}
                            )
                            if created:
                                added_count += 1
                            else:
                                updated_count += 1
                        except Exception:
                            continue
            elif ext == ".xls":
                import xlrd

                book = xlrd.open_workbook(file_contents=file.read())
                ws = book.sheet_by_index(0)
                for rx in range(ws.nrows):
                    if ws.ncols >= 2:
                        igk = str(ws.cell_value(rx, 1)).strip()
                        if not igk or igk.lower() in ["nan", "гк", "номер", "none"]:
                            continue
                        try:
                            rate_val = Decimal("22.0")
                            obj, created = GozContractVat.objects.update_or_create(
                                igk=igk, defaults={"vat_rate": rate_val}
                            )
                            if created:
                                added_count += 1
                            else:
                                updated_count += 1
                        except Exception:
                            continue
            else:
                messages.error(
                    request,
                    "Неподдерживаемый формат. Используйте .docx, .xlsx или .xls",
                )
                return redirect("goz_report")

        except Exception as e:
            messages.error(request, f"Ошибка парсинга файла: {e}")
            return redirect("goz_report")

        messages.success(
            request,
            f"Справочник обновлён. Добавлено: {added_count}, Обновлено: {updated_count}.",
        )
        return redirect("goz_report")

    return redirect("goz_report")


@login_required
def gk_directory_save(request):
    """Создание или редактирование записи справочника ГК."""
    if request.method == "POST":
        record_id = request.POST.get("record_id", "").strip()
        igk = request.POST.get("igk", "").strip()
        year = request.POST.get("year", "").strip()
        product = request.POST.get("product", "").strip()
        vat_rate_raw = request.POST.get("vat_rate", "22.0").strip()

        if not igk:
            messages.error(request, "Поле «ГК» обязательно для заполнения.")
            return redirect("goz_report")

        try:
            vat_rate = Decimal(vat_rate_raw.replace(",", "."))
        except Exception:
            vat_rate = Decimal("22.0")

        try:
            if record_id:
                obj = GozContractVat.objects.get(id=record_id)
                obj.igk = igk
                obj.year = year or None
                obj.product = product or None
                obj.vat_rate = vat_rate
                obj.save()
                messages.success(request, f"Запись для ГК «{igk}» обновлена.")
            else:
                GozContractVat.objects.create(
                    igk=igk,
                    year=year or None,
                    product=product or None,
                    vat_rate=vat_rate,
                )
                messages.success(request, f"ГК «{igk}» добавлен в справочник.")
        except Exception as e:
            messages.error(request, f"Ошибка сохранения: {e}")

    return redirect("goz_report")


@login_required
def gk_directory_delete(request):
    """Удаление записи справочника ГК."""
    if request.method == "POST":
        record_id = request.POST.get("record_id", "").strip()
        if record_id:
            try:
                obj = GozContractVat.objects.get(id=record_id)
                igk = obj.igk
                obj.delete()
                messages.success(request, f"Запись для ГК «{igk}» удалена.")
            except GozContractVat.DoesNotExist:
                messages.error(request, "Запись не найдена.")
            except Exception as e:
                messages.error(request, f"Ошибка удаления: {e}")
    return redirect("goz_report")


# --- Сводки ---


@login_required
def dashboard(request):
    """Сводка по договорам: плашки и таблица по ЦФО."""
    available_years = YEARS
    available_igk = NsiIgk.objects.all()
    year = valid_year(request.GET.get("year"))
    selected_igk = request.GET.get("igk", "") or str(available_igk.first() or "")

    ctx = _ctx(request)
    year_field = f"y{str(year)[-2:]}"
    concluded_q = Q(status__in=CONCLUDED)
    not_concl_q = Q(status__in=NOT_CONCL)
    year_q = Q(**{year_field: True})
    advance_q = Q(payment_type=ADVANCE)

    totals_all = IgkStatData.objects.filter(
        HAS_ORDER_Q, contract__isnull=False
    ).aggregate(
        count=Count("contract", distinct=True),
        plan_sum=Sum("plan"),
        concluded_count=Count("contract", filter=concluded_q, distinct=True),
        concluded_plan=Sum("plan", filter=concluded_q),
        not_concluded_count=Count("contract", filter=not_concl_q, distinct=True),
        not_concluded_plan=Sum("plan", filter=not_concl_q),
    )

    totals_year = (
        IgkStatData.objects.exclude(status="Расторгнут")
        .filter(HAS_ORDER_Q, contract__isnull=False, **{year_field: True})
        .aggregate(
            count=Count("contract", distinct=True),
            plan_sum=Sum("plan"),
            concluded_count=Count("contract", filter=concluded_q, distinct=True),
            concluded_plan=Sum("plan", filter=concluded_q),
            advance_plan=Sum("plan", filter=advance_q),
            advance_fact=Sum("fact", filter=advance_q),
        )
    )

    available_cfo = list(
        IgkStatData.objects.filter(igk=selected_igk)
        .values_list("cfo", flat=True)
        .distinct()
        .order_by("cfo")
    )

    cfo_stats = {
        row["cfo"]: row
        for row in (
            IgkStatData.objects.exclude(status="Расторгнут")
            .filter(HAS_ORDER_Q, igk=selected_igk, contract__isnull=False)
            .values("cfo")
            .annotate(
                all_count=Count("contract", distinct=True),
                all_sum=Sum("plan"),
                all_concluded_count=Count(
                    "contract", filter=concluded_q, distinct=True
                ),
                all_concluded_sum=Sum("plan", filter=concluded_q),
                curr_count=Count("contract", filter=year_q, distinct=True),
                curr_sum=Sum("plan", filter=year_q),
                curr_concluded_count=Count(
                    "contract", filter=concluded_q & year_q, distinct=True
                ),
                curr_concluded_sum=Sum("plan", filter=concluded_q & year_q),
                curr_not_concluded_count=Count(
                    "contract", filter=not_concl_q & year_q, distinct=True
                ),
                curr_not_concluded_sum=Sum("plan", filter=not_concl_q & year_q),
                curr_plan=Sum("plan", filter=advance_q & year_q),
                curr_fact=Sum("fact", filter=advance_q & year_q),
            )
        )
    }
    igk_table = [
        cfo_row(cfo, cfo_stats.get(cfo, EMPTY_CFO_STATS)) for cfo in available_cfo
    ]
    igk_table.append(cfo_totals_row(igk_table))

    year_count = totals_year["count"]
    year_concluded_count = totals_year["concluded_count"]
    year_plan = totals_year["plan_sum"] or 0
    year_concluded_plan = totals_year["concluded_plan"] or 0
    advance_plan = totals_year["advance_plan"] or 0
    advance_fact = totals_year["advance_fact"] or 0

    ctx.update(
        {
            "available_years": available_years,
            "selected_year": str(year),
            "available_igk": available_igk,
            "selected_igk": selected_igk,
            "all_contracts_count": totals_all["count"],
            "all_contracts_sum": to_mln(totals_all["plan_sum"]),
            "all_concluded_count": totals_all["concluded_count"],
            "all_concluded_sum": to_mln(totals_all["concluded_plan"]),
            "all_not_concluded_count": totals_all["not_concluded_count"],
            "all_not_concluded_sum": to_mln(totals_all["not_concluded_plan"]),
            "curr_year_contracts_count": year_count,
            "curr_year_contracts_sum": to_mln(year_plan),
            "curr_year_concluded_count": year_concluded_count,
            "curr_year_concluded_sum": to_mln(year_concluded_plan),
            "curr_year_not_concluded_count": year_count - year_concluded_count,
            "curr_year_not_concluded_sum": to_mln(year_plan - year_concluded_plan),
            "curr_year_fact": to_mln(advance_fact),
            "curr_year_plan": to_mln(advance_plan),
            "curr_year_percent_count": percent(year_concluded_count, year_count),
            "curr_year_percent_sum": percent(year_concluded_plan, year_plan),
            "curr_year_percent_prepaid": percent(advance_fact, advance_plan),
            "igk_table": igk_table,
            "has_data": IgkStatData.objects.exists(),
            "no_data_hint": "Договоры ещё не загружены. Нужен файл выгрузки по договорам.",
        }
    )
    return render(request, "dashboard.html", ctx)


@login_required
def znp_table(request):
    """Сводка заявок ФЗД: плашки, таблица по ЦФО, период для графика."""
    available_years = YEARS
    available_igk = NsiIgk.objects.all()
    year = valid_year(request.GET.get("year"))
    selected_igk = request.GET.get("igk", "") or str(available_igk.first() or "")

    ctx = _ctx(request)
    has_znp = Exists(ZnpData.objects.filter(parent=OuterRef("pk")))

    def _not_issued_qs(igk=None):
        """Заключённые позиции без заявок, где остаток превышает допуск."""
        qs = (
            IgkStatData.objects.filter(status__in=CONCLUDED)
            .annotate(has_znp=has_znp)
            .filter(has_znp=False)
            .annotate(needs_znp=RawSQL(needs_znp("igk_stat_data"), []))
            .filter(needs_znp=True)
        )
        if igk is not None:
            qs = qs.filter(igk=igk)
        return qs

    def _znp_qs(igk=None):
        """Заявки по заключённым позициям."""
        qs = ZnpData.objects.filter(parent__status__in=CONCLUDED)
        if igk is not None:
            qs = qs.filter(parent__igk=igk)
        return qs

    def _stages_qs(igk=None):
        """Все этапы оплаты заключённых позиций."""
        qs = IgkStatData.objects.filter(status__in=CONCLUDED)
        if igk is not None:
            qs = qs.filter(igk=igk)
        return qs

    all_not_issued_qs = _not_issued_qs()
    all_znp_qs = _znp_qs()
    all_stages_qs = _stages_qs()
    year_not_issued_qs = filter_by_year(all_not_issued_qs, year)
    year_znp_qs = filter_by_year(all_znp_qs, year, field_prefix="parent__")
    year_stages_qs = filter_by_year(all_stages_qs, year)

    def _breakdown(not_issued_qs, znp_qs, stages_qs):
        """Собирает карточки сводки из трёх групп агрегатов."""
        return breakdown_from_stats(
            not_issued_qs.aggregate(**not_issued_aggregates()),
            znp_qs.aggregate(**znp_aggregates()),
            stages_qs.aggregate(**stage_aggregates()),
        )

    all_breakdown = _breakdown(all_not_issued_qs, all_znp_qs, all_stages_qs)
    year_breakdown = _breakdown(year_not_issued_qs, year_znp_qs, year_stages_qs)

    available_cfo = list(
        IgkStatData.objects.filter(igk=selected_igk)
        .values_list("cfo", flat=True)
        .distinct()
        .order_by("cfo")
    )

    igk_not_issued_qs = filter_by_year(_not_issued_qs(igk=selected_igk), year)
    igk_znp_qs = filter_by_year(
        _znp_qs(igk=selected_igk), year, field_prefix="parent__"
    )
    igk_stages_qs = filter_by_year(_stages_qs(igk=selected_igk), year)

    not_issued_stats = {
        row["cfo"]: row
        for row in igk_not_issued_qs.values("cfo").annotate(**not_issued_aggregates())
    }
    znp_stats = {
        row["parent__cfo"]: row
        for row in igk_znp_qs.values("parent__cfo").annotate(**znp_aggregates())
    }
    stage_stats = {
        row["cfo"]: row
        for row in igk_stages_qs.values("cfo").annotate(**stage_aggregates())
    }

    cfo_table = [
        cfo_breakdown_row(
            cfo,
            breakdown_from_stats(
                not_issued_stats.get(cfo, EMPTY_NOT_ISSUED),
                znp_stats.get(cfo, EMPTY_ZNP),
                stage_stats.get(cfo, EMPTY_STAGES),
            ),
            ZNP_STAGE_LABELS,
        )
        for cfo in available_cfo
    ]
    cfo_total_row = cfo_breakdown_row(
        "ИТОГО",
        _breakdown(igk_not_issued_qs, igk_znp_qs, igk_stages_qs),
        ZNP_STAGE_LABELS,
    )

    chart_start = request.GET.get("start", "").strip()
    chart_end = request.GET.get("end", "").strip()
    if not (valid_date(chart_start) and valid_date(chart_end)):
        chart_start = chart_end = ""

    ctx.update(
        {
            "available_years": available_years,
            "selected_year": str(year),
            "available_igk": available_igk,
            "selected_igk": selected_igk,
            "chart_start": chart_start,
            "chart_end": chart_end,
            "stage_names": ZNP_STAGE_NAMES,
            "all": all_breakdown,
            "year": year_breakdown,
            "cfo_table": cfo_table,
            "cfo_total_row": cfo_total_row,
            "has_data": all_breakdown["total_count"] > 0,
            "no_data_hint": "Нет ни договоров, ни заявок на платёж (ФЗД). Нужны файлы выгрузки.",
        }
    )
    return render(request, "znp_table.html", ctx)


@login_required
def znp_sap_table(request):
    """Сводка заявок SAP: плашки, таблица по ЦФО, карточки по датам."""
    ctx = _ctx(request)
    qs = ZnpDataSAP.objects.annotate(sap_status=sap_status_expr()).filter(
        cfo__in=SAP_CFO
    )

    def _breakdown(qs):
        """Собирает карточки сводки: всего + по каждому статусу."""
        return sap_cards(
            qs.aggregate(total=Count("id"), total_sum=Sum("vv_sum")),
            {
                row["sap_status"]: row
                for row in qs.values("sap_status").annotate(
                    count=Count("id"), vv_sum=Sum("vv_sum")
                )
            },
        )

    all_breakdown = _breakdown(qs)

    date_param = request.GET.get("date", "")
    first_date = (
        datetime.strptime(date_param, "%Y-%m-%d").date()
        if valid_date(date_param)
        else timezone.localdate()
    )
    second_date = sap_second_date(first_date)
    first_date_breakdown = _breakdown(qs.filter(payment_possible=first_date))
    second_date_breakdown = _breakdown(qs.filter(payment_possible=second_date))

    available_igk = list(
        qs.exclude(igk__isnull=True)
        .exclude(igk="")
        .values_list("igk", flat=True)
        .distinct()
        .order_by("igk")
    )
    selected_igk = request.GET.get("igk", "") or (
        available_igk[0] if available_igk else ""
    )

    cfo_qs = qs.filter(igk=selected_igk) if selected_igk else qs
    available_cfo = list(
        cfo_qs.values_list("cfo", flat=True).distinct().order_by("cfo")
    )

    cfo_totals = {
        row["cfo"]: row
        for row in cfo_qs.values("cfo").annotate(
            total=Count("id"), total_sum=Sum("vv_sum")
        )
    }
    cfo_status = {}
    for row in cfo_qs.values("cfo", "sap_status").annotate(
        count=Count("id"), vv_sum=Sum("vv_sum")
    ):
        cfo_status.setdefault(row["cfo"], {})[row["sap_status"]] = row

    cfo_table = [
        cfo_breakdown_row(
            cfo,
            sap_cards(cfo_totals.get(cfo), cfo_status.get(cfo, {})),
            SAP_STAGE_PARAMS,
        )
        for cfo in available_cfo
    ]
    cfo_total_row = cfo_breakdown_row("ИТОГО", _breakdown(cfo_qs), SAP_STAGE_PARAMS)

    ctx.update(
        {
            "stage_names": SAP_STAGE_NAMES,
            "available_igk": available_igk,
            "selected_igk": selected_igk,
            "all": all_breakdown,
            "first_date": first_date,
            "second_date": second_date,
            "first_date_breakdown": first_date_breakdown,
            "second_date_breakdown": second_date_breakdown,
            "cfo_table": cfo_table,
            "cfo_total_row": cfo_total_row,
            "has_data": all_breakdown["total_count"] > 0,
            "no_data_hint": "Заявки на платёж из SAP ещё не загружены. Нужен файл выгрузки ЗНП (SAP).",
        }
    )

    try:
        sap_load_event = SystemEvent.objects.filter(event_key="sap_load").first()
        ctx["sap_load_time"] = sap_load_event.event_time if sap_load_event else None
    except Exception:
        ctx["sap_load_time"] = None
    return render(request, "znp_sap_table.html", ctx)
