"""
Анализ отчётов ЕИС ГОЗ (форма исполнения госконтракта) из ZIP-архива .xls файлов.

Один .xls = один контракт (имя файла = последние 4 цифры ИГК).
Расположение показателей на листе «Лист_1» фиксировано (форма госсистемы):
колонка D — целевые (план), колонка G — сальдо операций (факт).
"""

import zipfile
from io import BytesIO

import xlrd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --- Константы для парсинга .xls ---

SHEET_NAME = "Лист_1"
_PLAN_COL = 3  # столбец D
_FACT_COL = 6  # столбец G

# Строки формы, присутствующие в плане и факте
_LINE_ROWS = {
    "shipment": 31,  # 3. Отгрузка товара, выполнение работ, оказание услуг
    "cost": 32,  # 3.1 Себестоимость реализованной продукции
    "amr": 33,  # 3.2 Административно-управленческие расходы
    "commercial": 34,  # 3.3 Коммерческие расходы
    "credit_pct": 35,  # 3.4 Проценты по кредитам банка
    "profit": 37,  # 3.6 Прибыль контракта
}

# Строки формы, присутствующие только в факте (сальдо операций)
_FACT_ONLY_ROWS = {
    "financing": 4,  # 1. Финансирование контракта
    "distribution": 9,  # 2. Распределение ресурсов контракта
    "materials": 15,  # 2.2.1 Материалы на складах (ТМЦ)
    "vat_in": 16,  # 2.2.2 НДС входящий
    "wip": 21,  # 2.3 Производство (НЗП)
    "resource_delta": 38,  # 4. (+/-) Привлечение/перенаправление ресурсов
}

# --- Константы для генерации .xlsx ---

HEADERS = [
    "№ п/п",
    "ГК",
    "Наименование",
    "",
    "Отгрузка (п.3)",
    "Себестоимость (п.3.1)",
    "АУР (п.3.2)",
    "КР (п.3.3)",
    "НДС",
    "% кредит (п.3.4)",
    "Прибыль в отчёте (п.3.6)",
    "Прибыль (проверка)",
    "Рент. к с/с, % (проверка)",
    "Рент. к с/с, % по прибыли в отчёте",
    "Финансирование контракта (п.1)",
    "Распределение ресурсов (п.2)",
    "+/- ресурсов ГК (проверка)",
    "+/- ресурсов ГК (п.4)",
    "НЗП (п.2.3)",
    "ТМЦ (п.2.2.1)",
    "НДС вх. (п.2.2.2)",
    "Примечание",
    "Комментарий",
]
TOTAL_COLS = len(HEADERS)  # 23

# Ширина колонок: №, ГК, Наименование, метка строки, E..N (деньги/деньги/.../%/%),
# O..U (деньги), Примечание, Комментарий. Денежные — широкие, чтобы не было "#####".
_COL_WIDTHS = (
    [6, 8, 16, 12] + [19, 19, 19, 19, 19, 19, 19, 19, 10, 10] + [19] * 7 + [30, 30]
)

# Колонки, объединяемые на 4 строки блока (значение одно на весь ГК).
_MERGE_COLS = [1, 2, 3, 22, 23]  # №, ГК, Наименование, Примечание, Комментарий

_K_COL, _L_COL, _M_COL, _N_COL = 11, 12, 13, 14
_Q_COL, _R_COL = 17, 18
_PERCENT_COLS = {_M_COL, _N_COL}

_RENT_THRESHOLD = 0.075  # порог рентабельности 7.5%
_MISMATCH_EPS = 1.0  # допуск в рублях при сравнении

# --- Стили ячеек ---

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_FONT_NAME = "Times New Roman"
_FONT = Font(name=_FONT_NAME, size=10)
_TITLE_FONT = Font(name=_FONT_NAME, size=14, bold=True)
_RED_FONT = Font(name=_FONT_NAME, size=10, color="9C0006")
_RENT_FILL = PatternFill(
    "solid", fgColor="FFC7CE"
)  # рентабельность > 7.5% (с красным шрифтом)
_NEGATIVE_FILL = PatternFill("solid", fgColor="FFC7CE")  # Q < 0 (без красного шрифта)
_PROFIT_MISMATCH_FILL = PatternFill("solid", fgColor="CCC1DA")  # K ≠ L
_RESOURCE_MISMATCH_FILL = PatternFill("solid", fgColor="FAC090")  # Q ≠ R
_BOTTOM_ONLY = Border(bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
_RIGHT = Alignment(horizontal="right")

# Пункты легенды: (заливка, шрифт-акцент, текст пояснения) — цвета и охват
# в точности как в исходных правилах условного форматирования файла-примера.
_LEGEND_ITEMS = [
    (_RENT_FILL, _RED_FONT, "Рентабельность к себестоимости (столбцы M, N) выше 7,5%"),
    (_NEGATIVE_FILL, _FONT, "«+/- ресурсов ГК» (столбец Q) — отрицательное значение"),
    (
        _PROFIT_MISMATCH_FILL,
        _FONT,
        "Прибыль в отчёте (K) не совпадает с расчётной прибылью (L)",
    ),
    (
        _RESOURCE_MISMATCH_FILL,
        _FONT,
        "«+/- ресурсов ГК»: расчёт (Q) не совпадает со значением из формы (R)",
    ),
]


# --- Парсинг .xls ---


def _cell_number(sheet, row, col):
    """Возвращает число из ячейки или 0.0 для нечисловых значений."""
    try:
        value = sheet.cell_value(row, col)
    except IndexError:
        return 0.0
    return float(value) if isinstance(value, (int, float)) else 0.0


def _read_contract_file(igk, content):
    """Парсит один .xls файл и возвращает (igk, plan, fact)."""
    try:
        book = xlrd.open_workbook(file_contents=content)
        sheet = book.sheet_by_name(SHEET_NAME)
    except Exception as e:
        raise ValueError(f"файл {igk}.xls: {e}") from e

    plan = {k: _cell_number(sheet, r, _PLAN_COL) for k, r in _LINE_ROWS.items()}
    fact = {k: _cell_number(sheet, r, _FACT_COL) for k, r in _LINE_ROWS.items()}
    fact.update(
        {k: _cell_number(sheet, r, _FACT_COL) for k, r in _FACT_ONLY_ROWS.items()}
    )
    return igk, plan, fact


def read_archive(archive_file):
    """Читает ZIP-архив с .xls отчётами. Возвращает отсортированный список (igk, plan, fact)."""
    contracts = []
    with zipfile.ZipFile(archive_file) as zf:
        for name in zf.namelist():
            if name.endswith("/") or not name.lower().endswith(".xls"):
                continue
            igk = name.rsplit("/", 1)[-1][:-4]
            with zf.open(name) as fh:
                contracts.append(_read_contract_file(igk, fh.read()))

    if not contracts:
        raise ValueError("В архиве не найдено ни одного .xls файла")

    contracts.sort(key=lambda c: (len(c[0]), c[0]))
    return contracts


# --- Расчёт метрик ---


def _line_metrics(values, vat_rate):
    """
    Вычисляет 10 показателей (колонки E..N) из 6 исходных строк формы.
    Возвращает [shipment, cost, amr, commercial, vat, credit_pct, profit, calc_profit, rent_calc, rent_report].
    """
    shipment = values["shipment"]
    cost = values["cost"]
    amr = values["amr"]
    commercial = values["commercial"]
    vat = shipment - shipment / (1 + vat_rate / 100)
    credit_pct = values["credit_pct"]
    profit = values["profit"]
    calc_profit = shipment - cost - amr - commercial - vat - credit_pct
    rent_calc = (calc_profit / cost) if cost else None
    rent_report = (profit / cost) if cost else None
    return [
        shipment,
        cost,
        amr,
        commercial,
        vat,
        credit_pct,
        profit,
        calc_profit,
        rent_calc,
        rent_report,
    ]


def _deviation(fact_m, plan_m):
    """Вычисляет абсолютные и относительные отклонения факта от плана."""
    abs_dev = [
        (
            (fv - pv)
            if isinstance(fv, (int, float)) and isinstance(pv, (int, float))
            else None
        )
        for fv, pv in zip(fact_m, plan_m)
    ]
    rel_dev = [
        (av / pv) if av is not None and pv else None for av, pv in zip(abs_dev, plan_m)
    ]
    return abs_dev, rel_dev


def _block_rows(igk, num, plan, fact, vat_rate):
    """
    Строит 4 строки блока (Целевые/Факт/откл.абсол./откл.отн.%), каждая
    ровно TOTAL_COLS элементов — чтобы рамка потом легла на всю таблицу
    целиком, а не только на колонки с данными.
    """
    plan_m = _line_metrics(plan, vat_rate)
    fact_m = _line_metrics(fact, vat_rate)
    fact_extra = [
        fact["financing"],
        fact["distribution"],
        fact["distribution"] - fact["financing"],
        fact["resource_delta"],
        fact["wip"],
        fact["materials"],
        fact["vat_in"],
    ]
    abs_dev, rel_dev = _deviation(fact_m, plan_m)
    pad_extra = [None] * 7  # для строк, где O..U не считаются
    pad_notes = [None, None]  # Примечание, Комментарий — всегда пустые

    return [
        [num, igk, "", "Целевые", *plan_m, *pad_extra, *pad_notes],
        [None, None, None, "Факт", *fact_m, *fact_extra, *pad_notes],
        [None, None, None, "откл.абсол.", *abs_dev, *pad_extra, *pad_notes],
        [None, None, None, "откл.отн.%", *rel_dev, *pad_extra, *pad_notes],
    ]


# --- Запись строк в .xlsx ---


def _write_row(ws, row_idx, values, force_percent=False):
    """Записывает строку целиком по TOTAL_COLS колонок — рамка на каждой ячейке."""
    for ci, val in enumerate(values, 1):
        cell = ws.cell(row=row_idx, column=ci, value=val)
        cell.border = _BORDER
        cell.font = _FONT

        is_num = isinstance(val, (int, float)) and not isinstance(val, bool)
        if ci in (1, 2):
            cell.alignment = _CENTER
        elif ci in (3, 22, 23):
            cell.alignment = _LEFT_WRAP
        elif is_num and ci >= 5:
            is_percent = force_percent or ci in _PERCENT_COLS
            cell.number_format = "0.0%" if is_percent else "#,##0.00"
            cell.alignment = _RIGHT


def _apply_report_checks(ws, row_idx, metrics):
    """
    Подсветка для строк «Целевые» и «Факт» (цвета — как в примере):
    - светло-сиреневый (CCC1DA): K (прибыль в отчёте) ≠ L (прибыль расчётная)
    - красный (FFC7CE) + красный шрифт: M/N (рентабельность) > 7.5%
    """
    profit_report, profit_calc = metrics[6], metrics[7]
    if isinstance(profit_report, (int, float)) and isinstance(
        profit_calc, (int, float)
    ):
        if abs(profit_report - profit_calc) > _MISMATCH_EPS:
            ws.cell(row=row_idx, column=_K_COL).fill = _PROFIT_MISMATCH_FILL
            ws.cell(row=row_idx, column=_L_COL).fill = _PROFIT_MISMATCH_FILL

    for val, col in ((metrics[8], _M_COL), (metrics[9], _N_COL)):
        if isinstance(val, (int, float)) and val > _RENT_THRESHOLD:
            cell = ws.cell(row=row_idx, column=col)
            cell.fill = _RENT_FILL
            cell.font = _RED_FONT


def _apply_fact_checks(ws, row_idx, fact_extra):
    """
    Подсветка только для строки «Факт» (цвета — как в примере):
    - светло-оранжевый (FAC090): Q (+/- ресурсов расчёт) ≠ R (значение из формы)
    - красный (FFC7CE), без красного шрифта: Q < 0 (перекрывает предыдущую заливку)
    """
    q_calc, r_report = fact_extra[2], fact_extra[3]
    if isinstance(q_calc, (int, float)) and isinstance(r_report, (int, float)):
        if abs(q_calc - r_report) > _MISMATCH_EPS:
            ws.cell(row=row_idx, column=_Q_COL).fill = _RESOURCE_MISMATCH_FILL
            ws.cell(row=row_idx, column=_R_COL).fill = _RESOURCE_MISMATCH_FILL
    if isinstance(q_calc, (int, float)) and q_calc < 0:
        ws.cell(row=row_idx, column=_Q_COL).fill = _NEGATIVE_FILL


def _write_legend(ws, start_row):
    """Пишет легенду цветовых отметок: цветной образец + пояснение в строке."""
    header = ws.cell(row=start_row, column=1, value="Легенда цветовых отметок:")
    header.font = Font(name=_FONT_NAME, size=11, bold=True)

    for i, (fill, font, text) in enumerate(_LEGEND_ITEMS):
        r = start_row + 1 + i
        swatch = ws.cell(row=r, column=1)
        swatch.fill = fill
        swatch.border = _BORDER
        label = ws.cell(row=r, column=2, value=text)
        label.font = font
        label.alignment = Alignment(horizontal="left", vertical="center")
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=8)


def _merge_block(ws, first_row):
    """Объединяет №/ГК/Наименование/Примечание/Комментарий на все 4 строки блока."""
    last_row = first_row + 3
    for col in _MERGE_COLS:
        ws.merge_cells(
            start_row=first_row, start_column=col, end_row=last_row, end_column=col
        )


# --- Генерация отчёта ---


def build_report(contracts, vat_rates):
    """
    Генерирует Excel-отчёт анализа ГОЗ.

    Args:
        contracts: результат read_archive()
        vat_rates: dict {igk: vat_rate} или float для совместимости

    Returns:
        bytes готового .xlsx файла
    """
    if isinstance(vat_rates, (int, float, str)):
        # fallback для совместимости - если передана одна ставка
        try:
            rate = float(str(vat_rates).replace(",", "."))
        except:
            rate = 20.0
        vat_rates_dict = {c[0]: rate for c in contracts}
    else:
        vat_rates_dict = {}
        for k, v in vat_rates.items():
            try:
                vat_rates_dict[k] = float(str(v).replace(",", "."))
            except:
                vat_rates_dict[k] = 20.0

    wb = Workbook()
    ws = wb.active
    ws.title = "Анализ ГОЗ"
    ws.freeze_panes = "A3"

    title_cell = ws.cell(row=1, column=1, value="Анализ отчётов ЕИС ГОЗ")
    title_cell.font = _TITLE_FONT
    title_cell.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TOTAL_COLS)
    for ci in range(1, TOTAL_COLS + 1):
        ws.cell(row=1, column=ci).border = _BOTTOM_ONLY
    ws.row_dimensions[1].height = 24

    hdr_aln = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for ci, (h, w) in enumerate(zip(HEADERS, _COL_WIDTHS), 1):
        cell = ws.cell(row=2, column=ci, value=h)
        cell.font = _FONT
        cell.border = _BORDER
        cell.alignment = hdr_aln
        ws.column_dimensions[get_column_letter(ci)].width = w
    ws.row_dimensions[2].height = 60

    row_i = 3
    for num, (igk, plan, fact) in enumerate(contracts, 1):
        # Берем индивидуальную ставку для каждого ГК
        rate = vat_rates_dict.get(igk, 20.0)
        block = _block_rows(igk, num, plan, fact, rate)

        _write_row(ws, row_i, block[0])
        plan_m = _line_metrics(plan, rate)
        _apply_report_checks(ws, row_i, plan_m)
        row_i += 1

        _write_row(ws, row_i, block[1])
        fact_m = _line_metrics(fact, rate)
        fact_extra = block[1][4 + 10 : 4 + 10 + 7]
        _apply_report_checks(ws, row_i, fact_m)
        _apply_fact_checks(ws, row_i, fact_extra)
        row_i += 1

        _write_row(ws, row_i, block[2])
        row_i += 1

        _write_row(ws, row_i, block[3], force_percent=True)
        row_i += 1

        _merge_block(ws, row_i - 4)

    _write_legend(ws, row_i + 1)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
