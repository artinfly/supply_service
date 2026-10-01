"""
Формирование Краткой справки (обычная и длинноцикловая) по шаблону.

Логика перенесена из Delphi-процедур SumReport / SumReportWCycle.
Для каждого ИГК (листа исходного файла) генерируется отдельный xlsx,
все файлы упаковываются в ZIP-архив.

ВАЖНО: шаблон должен лежать в reports/static/files/templateSum.xlsx.
Шаблон содержит заранее заложенные итоговые строки для каждой таблицы.
При вставке строк данных эти итоговые строки сдвигаются вниз,
и мы записываем суммы именно в них.
"""

import io
import os
import zipfile
from copy import copy
from decimal import Decimal

from django.db import connection
from django.utils import timezone
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from .excel import xlsx_response

TEMPLATE_NAME = "templateSum.xlsx"

# Красная заливка для подсветки расхождений (аналог ColorIndex := 3 в Delphi)
RED_FILL = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")


def _template_path():
    """Полный путь к шаблону отчёта."""
    return os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "static",
        "files",
        TEMPLATE_NAME,
    )


def _copy_cell_style(source_cell, target_cell):
    """Копирует стиль ячейки (шрифт, границы, заливку, формат, выравнивание)."""
    if source_cell.has_style:
        target_cell.font = copy(source_cell.font)
        target_cell.border = copy(source_cell.border)
        target_cell.fill = copy(source_cell.fill)
        target_cell.number_format = source_cell.number_format
        target_cell.protection = copy(source_cell.protection)
        target_cell.alignment = copy(source_cell.alignment)


def _to_float(v):
    """Приводит значение к float, None -> 0.0."""
    if v is None:
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0.0


# ---------- SQL-запросы ----------


def _sql_table1(igk, is_cycle):
    """Первая таблица: сводная по ЦФО."""
    cycle_cond = "AND is_cycle = TRUE" if is_cycle else ""
    cycle_sub = "AND td2.is_cycle = TRUE" if is_cycle else ""
    return f""" 
        SELECT 
            td.dep, 
            COUNT(*) AS contractCount, 
            ROUND(SUM(td.contract_sum)::numeric, 2) AS contractSum, 
            COALESCE(( 
                SELECT COUNT(*) FROM staging_sum_excel td2 
                WHERE td.dep = td2.dep AND td2.igk = %s {cycle_sub} 
                  AND LOWER(td2.status) IN ('заключен', ' заключен') 
                GROUP BY td2.dep 
            ), 0) AS completedCount, 
            COALESCE(( 
                SELECT ROUND(SUM(td2.contract_sum)::numeric, 2) FROM staging_sum_excel td2 
                WHERE td.dep = td2.dep AND td2.igk = %s {cycle_sub} 
                  AND LOWER(td2.status) IN ('заключен', ' заключен') 
                GROUP BY td2.dep 
            ), 0) AS completedSum, 
            COALESCE(( 
                SELECT ROUND(SUM(td2.sum_80)::numeric, 2) FROM staging_sum_excel td2 
                WHERE td.dep = td2.dep AND td2.igk = %s {cycle_sub} 
                GROUP BY td2.dep 
            ), 0) AS planAvans, 
            COALESCE(( 
                SELECT COUNT(*) FROM staging_sum_excel td2 
                WHERE td.dep = td2.dep AND td2.igk = %s {cycle_sub} 
                  AND td2.znp_count <> 0 
                GROUP BY td2.dep 
            ), 0) AS znpCount, 
            COALESCE(( 
                SELECT ROUND(SUM(td2.completed_sum)::numeric, 2) FROM staging_sum_excel td2 
                WHERE td.dep = td2.dep AND td2.igk = %s {cycle_sub} 
                GROUP BY td2.dep 
            ), 0) AS znpSum, 
            COALESCE(( 
                SELECT ROUND(SUM(td2.paid_from_znp)::numeric, 2) FROM staging_sum_excel td2 
                WHERE td.dep = td2.dep AND td2.igk = %s {cycle_sub} 
                GROUP BY td2.dep 
            ), 0) AS znpPaidSum, 
            COALESCE(( 
                SELECT ROUND(ABS(SUM(td2.sum_avans))::numeric, 2) FROM staging_sum_excel td2 
                WHERE td.dep = td2.dep AND td2.igk = %s {cycle_sub} 
                  AND td2.znp_count <> 0 AND td2.remains_pay <> 0 
                GROUP BY td2.dep 
            ), 0) AS remainPay, 
            COALESCE(( 
                SELECT ROUND(SUM(td2.remains_pay)::numeric, 2) FROM staging_sum_excel td2 
                WHERE td.dep = td2.dep AND td2.igk = %s {cycle_sub} 
                GROUP BY td2.dep 
            ), 0) AS totalRemainPay 
        FROM staging_sum_excel td 
        WHERE td.igk = %s {cycle_cond} 
          AND UPPER(TRIM(td.status)) IN ('ЗАКЛЮЧЕН', 'НЕ ЗАКЛЮЧЕН') 
        GROUP BY td.dep 
        ORDER BY td.dep ASC; 
    """


def _sql_table2(igk, is_cycle):
    """Вторая таблица: ЗнП оформлена / Заключён."""
    cycle_cond = "AND td.is_cycle = TRUE" if is_cycle else ""
    return f""" 
        SELECT 
            td.counteragent AS counteragent, 
            td.dep AS cfo, 
            td.item AS contractObject, 
            ROUND(td.contract_sum::numeric, 2) AS contractSum, 
            ROUND(ABS(td.sum_avans)::numeric, 2) AS sumAvans, 
            td.percent_doc * 100 AS percentDoc, 
            td.status AS condition, 
            td.period_reg_date AS periodRegDate, 
            td.note AS note 
        FROM staging_sum_excel td 
        WHERE td.igk = %s {cycle_cond} 
          AND LOWER(td.status) IN ('знп оформлена', 'заключен', ' заключен') 
          AND td.remains_pay NOT BETWEEN -1 AND 1 
          AND td.sum_avans NOT BETWEEN -1 AND 1 
        ORDER BY td.dep ASC; 
    """


def _sql_table3(igk, is_cycle):
    """Третья таблица: ЗнП на оформлении / Заключён."""
    cycle_cond = "AND td.is_cycle = TRUE" if is_cycle else ""
    return f""" 
        SELECT 
            td.counteragent AS counteragent, 
            td.dep AS cfo, 
            td.item AS contractObject, 
            ROUND(td.contract_sum::numeric, 2) AS contractSum, 
            ROUND(ABS(td.sum_issued_znp)::numeric, 2) AS sumAvans, 
            td.percent_doc * 100 AS percentDoc, 
            td.status AS condition, 
            td.period_reg_date AS periodRegDate, 
            td.note AS note 
        FROM staging_sum_excel td 
        WHERE td.igk = %s {cycle_cond} 
          AND LOWER(td.status) IN ('знп на оформлении', 'заключен', ' заключен', 'заключен ', ' заключен ') 
          AND td.remains_pay NOT BETWEEN -1 AND 1 
          AND ROUND(td.sum_issued_znp::numeric, 2) NOT BETWEEN -1 AND 1 
        ORDER BY td.dep ASC; 
    """


def _sql_table4(igk, is_cycle):
    """Четвёртая таблица: Не заключён."""
    cycle_cond = "AND td.is_cycle = TRUE" if is_cycle else ""
    return f""" 
        SELECT 
            td.counteragent AS counteragent, 
            td.dep AS cfo, 
            td.item AS contractObject, 
            ROUND(td.contract_sum::numeric, 2) AS contractSum, 
            ROUND(ABS(td.remains_pay)::numeric, 2) AS sumAvans, 
            td.percent_doc * 100 AS percentDoc, 
            CASE WHEN td.plan_date_contract IS NULL THEN '' ELSE TO_CHAR(td.plan_date_contract, 'DD.MM.YYYY') END AS planDateContract, 
            CASE WHEN td.period_reg_date IS NULL THEN '' ELSE TO_CHAR(td.period_reg_date, 'DD.MM.YYYY') END AS periodRegDate, 
            td.note AS note 
        FROM staging_sum_excel td 
        WHERE td.igk = %s {cycle_cond} 
          AND LOWER(td.status) IN ('не заключен', ' не заключен', 'не заключен ', ' не заключен ') 
        ORDER BY td.dep ASC; 
    """


def _fetch_rows(sql, params):
    """Выполняет SQL и возвращает список словарей."""
    with connection.cursor() as cur:
        cur.execute(sql, params)
        cols = [c[0] for c in cur.description]
        return [dict(zip(cols, r)) for r in cur.fetchall()]


def _get_igk_list(is_cycle):
    """Список ИГК, для которых есть данные в staging."""
    with connection.cursor() as cur:
        cur.execute(
            "SELECT DISTINCT igk FROM staging_sum_excel WHERE is_cycle = %s ORDER BY igk",
            [is_cycle],
        )
        return [row[0] for row in cur.fetchall()]


def _write_table1(ws, igk, is_cycle):
    """
    Заполняет первую таблицу (сводная по ЦФО).
    Возвращает номер итоговой строки первой таблицы.
    """
    rows = _fetch_rows(_sql_table1(igk, is_cycle), [igk] * 9)
    if not rows:
        return 5

    # Сохраняем стиль первой строки данных (строка 5 в шаблоне)
    template_styles = [ws.cell(row=5, column=c) for c in range(1, 12)]
    curr_row = 5

    for row_data in rows:
        if not row_data.get("dep"):
            continue
        ws.insert_rows(curr_row)
        values = [
            row_data["dep"],
            row_data["contractcount"],
            row_data["contractsum"],
            row_data["completedcount"],
            row_data["completedsum"],
            row_data["planavans"],
            row_data["znpcount"],
            row_data["znpsum"],
            row_data["znppaidsum"],
            row_data["remainpay"],
            row_data["totalremainpay"],
        ]
        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(row=curr_row, column=c_idx)
            _copy_cell_style(template_styles[c_idx - 1], cell)
            cell.value = val
        curr_row += 1

    # Итоговая строка уже существует в шаблоне (сдвинута вниз).
    # Считаем суммы по всем вставленным строкам и записываем в итог.
    total_row = curr_row
    count_res_rows = len([r for r in rows if r.get("dep")])
    for offset in range(1, count_res_rows + 1):
        src_row = total_row - offset
        for c_idx in range(2, 12):
            dst_cell = ws.cell(row=total_row, column=c_idx)
            src_cell = ws.cell(row=src_row, column=c_idx)
            dst_cell.value = _to_float(dst_cell.value) + _to_float(src_cell.value)

    return total_row


def _write_detail_table(ws, curr_row, rows, is_table4=False):
    """
    Заполняет детальную таблицу (2, 3 или 4).
    Возвращает (новый_ряд, сумма_договоров, сумма_авансов, ряд_итога_таблицы).

    В шаблоне итоговая строка таблицы уже существует и сдвигается вниз.
    Строки с суммой договора < 1 000 000 накапливаются в «Прочие контрагенты».
    """
    # Сохраняем стиль строки данных (первая строка области данных таблицы)
    template_styles = [ws.cell(row=curr_row, column=c) for c in range(1, 10)]

    another_contract_sum = 0.0
    another_avans_sum = 0.0
    inserted_count = 0

    for row_data in rows:
        contract_sum = _to_float(row_data.get("contractsum"))
        sum_avans = _to_float(row_data.get("sumavans"))

        if contract_sum < 1000000:
            another_contract_sum += contract_sum
            another_avans_sum += sum_avans
            continue

        ws.insert_rows(curr_row)
        if is_table4:
            values = [
                row_data.get("counteragent"),
                row_data.get("contractobject"),
                contract_sum,
                sum_avans,
                f"{round(sum_avans / contract_sum * 100)}%" if contract_sum else "0%",
                row_data.get("plandatecontract"),
                row_data.get("periodregdate"),
                row_data.get("note"),
                row_data.get("cfo"),
            ]
        else:
            values = [
                row_data.get("counteragent"),
                row_data.get("contractobject"),
                contract_sum,
                sum_avans,
                f"{round(sum_avans / contract_sum * 100)}%" if contract_sum else "0%",
                row_data.get("condition"),
                row_data.get("periodregdate"),
                row_data.get("note"),
                row_data.get("cfo"),
            ]
        for c_idx, val in enumerate(values, 1):
            cell = ws.cell(row=curr_row, column=c_idx)
            _copy_cell_style(template_styles[c_idx - 1], cell)
            cell.value = val
        curr_row += 1
        inserted_count += 1

    # Итоговая строка таблицы (сдвинута вниз)
    total_row = curr_row

    # Суммируем вставленные строки в итоговую строку таблицы
    for offset in range(1, inserted_count + 1):
        src_row = total_row - offset
        for c_idx in (3, 4):
            dst_cell = ws.cell(row=total_row, column=c_idx)
            src_cell = ws.cell(row=src_row, column=c_idx)
            dst_cell.value = _to_float(dst_cell.value) + _to_float(src_cell.value)

    # Вставляем строку «Прочие контрагенты» перед итоговой строкой
    ws.insert_rows(total_row)
    ws.cell(row=total_row, column=1).value = "Прочие контрагенты:"
    ws.cell(row=total_row, column=3).value = another_contract_sum
    ws.cell(row=total_row, column=4).value = another_avans_sum

    # Прибавляем «прочих» к итоговой строке таблицы (она сдвинулась на 1 вниз)
    final_total_row = total_row + 1
    ws.cell(row=final_total_row, column=3).value = (
        _to_float(ws.cell(row=final_total_row, column=3).value) + another_contract_sum
    )
    ws.cell(row=final_total_row, column=4).value = (
        _to_float(ws.cell(row=final_total_row, column=4).value) + another_avans_sum
    )

    contract_sum_total = _to_float(ws.cell(row=final_total_row, column=3).value)
    avans_sum_total = _to_float(ws.cell(row=final_total_row, column=4).value)

    # После вставки «прочих» текущий ряд сдвигается на 1 вниз
    new_curr_row = total_row + 1
    return new_curr_row, contract_sum_total, avans_sum_total, final_total_row


def _generate_single_report(igk, is_cycle):
    """Генерирует один отчёт для одного ИГК. Возвращает bytes xlsx."""
    template = _template_path()
    if not os.path.exists(template):
        raise FileNotFoundError(f"Шаблон {template} не найден")

    wb = load_workbook(template)
    ws = wb.active

    # Заголовок: в оригинале брали значение из исходного файла [1,2],
    # у нас имя листа = ИГК, поэтому пишем ИГК.
    ws["A2"] = igk
    # Дата отчёта в [1,11]
    ws.cell(row=1, column=11).value = timezone.localdate().strftime("%d.%m.%Y")

    # --- Первая таблица ---
    fst_table_last_row = _write_table1(ws, igk, is_cycle)

    # --- Вторая таблица ---
    rows2 = _fetch_rows(_sql_table2(igk, is_cycle), [igk])
    curr_row = fst_table_last_row + 4
    curr_row, contract_sum_total, avans_sum_total, sec_table_last_row = (
        _write_detail_table(ws, curr_row, rows2, is_table4=False)
    )

    # --- Третья таблица ---
    rows3 = _fetch_rows(_sql_table3(igk, is_cycle), [igk])
    curr_row = curr_row + 3
    curr_row, c3, a3, _ = _write_detail_table(ws, curr_row, rows3, is_table4=False)
    contract_sum_total += c3
    avans_sum_total += a3

    # --- Четвёртая таблица ---
    rows4 = _fetch_rows(_sql_table4(igk, is_cycle), [igk])
    curr_row = curr_row + 3
    curr_row, c4, a4, _ = _write_detail_table(ws, curr_row, rows4, is_table4=True)
    contract_sum_total += c4
    avans_sum_total += a4

    # Общие суммы всего отчёта
    ws.cell(row=curr_row + 1, column=3).value = contract_sum_total
    ws.cell(row=curr_row + 1, column=4).value = avans_sum_total

    # Подстановка имени ИГК в ячейку с плейсхолдером __g
    cell_a = ws.cell(row=curr_row + 1, column=1)
    if cell_a.value and isinstance(cell_a.value, str):
        cell_a.value = cell_a.value.replace("__g", igk)

    # Переименование листа
    ws.title = igk

    # Проверка расхождений и подсветка красным
    val10 = _to_float(ws.cell(row=fst_table_last_row, column=10).value)
    val_sec4 = _to_float(ws.cell(row=sec_table_last_row, column=4).value)
    if abs(val10 - val_sec4) > 1:
        ws.cell(row=fst_table_last_row, column=10).fill = RED_FILL

    val11 = _to_float(ws.cell(row=fst_table_last_row, column=11).value)
    val_last4 = _to_float(ws.cell(row=curr_row + 1, column=4).value)
    if abs(val11 - val_last4) > 1:
        ws.cell(row=fst_table_last_row, column=11).fill = RED_FILL

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def generate_sum_reports_zip(is_cycle=False):
    """
    Формирует ZIP-архив с отчётами по всем ИГК.
    is_cycle: False = обычная Краткая справка, True = Длинноцикловая.
    """
    igks = _get_igk_list(is_cycle)
    if not igks:
        raise ValueError(
            "Нет данных для формирования отчёта. "
            "Загрузите файл Краткой справки на странице загрузки."
        )

    report_label = "Краткая (Длинноцикл.)" if is_cycle else "Краткая"
    today_str = timezone.localdate().strftime("%d-%m-%Y")
    username = "user"

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for igk in igks:
            try:
                data = _generate_single_report(igk, is_cycle)
            except Exception as exc:
                raise RuntimeError(f"Ошибка формирования отчёта ИГК {igk}: {exc}") from exc
            filename = f"{igk} {report_label} {today_str}_{username}.xlsx"
            zf.writestr(filename, data)

    zip_buffer.seek(0)
    from django.http import HttpResponse
    response = HttpResponse(zip_buffer.getvalue(), content_type="application/zip")
    response["Content-Disposition"] = f'attachment; filename="reports.zip"'
    return response
